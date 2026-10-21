"""
직무 스킬 분석 엑셀 리포트(2026-10-21) — job_skill_analysis.analyze() 결과를 시트별로 쓴다.
실행: scripts/analyze_job_skills.py. 값만 쓰는 정적 리포트(수식 없음).
"""
from __future__ import annotations

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = '맑은 고딕'
NAVY = '0B1B33'
F_BASE = Font(name=FONT, size=10)
F_BOLD = Font(name=FONT, size=10, bold=True)
F_HEAD = Font(name=FONT, size=10, bold=True, color='FFFFFF')
F_TITLE = Font(name=FONT, size=16, bold=True, color=NAVY)
F_SUB = Font(name=FONT, size=12, bold=True, color=NAVY)
F_NOTE = Font(name=FONT, size=9, color='5B6577')
FILL_HEAD = PatternFill('solid', fgColor=NAVY)
FILL_BLOCK = PatternFill('solid', fgColor='E8F0FE')
FILL_ALT = PatternFill('solid', fgColor='F6F8FB')
THIN = Side(style='thin', color='D5DBE5')
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical='top')
CENTER = Alignment(horizontal='center', vertical='center', wrap_text=True)
JUDGE_FILL = {'같은 내용': 'DFF3E8', '일부 다름': 'FFF4D6', '내용 다름': 'FDE2E1'}
KIND_FILL = {'공통': 'DFF3E8', '일부': 'E8F0FE', '고유': 'FFF1E3', '단일 부서': 'F1F3F7'}

LONG_TEXT = {'스킬 정의', 'Level 1', 'Level 2', 'Level 3', '정의 A', '정의 B', '점검 결과', '내용', '직무', '부서',
             '직무 A', '직무 B', '부서 A', '부서 B', '가장 다른 부서 쌍'}
SIM_WORDS = ('유사도', '일치도', '일치율')


def _width(col: str, values) -> float:
    if col in {'스킬 정의', 'Level 1', 'Level 2', 'Level 3', '정의 A', '정의 B'}:
        return 42
    longest = max([len(str(col))] + [len(str(v)) for v in list(values)[:300]])
    return max(8, min(36, longest * 1.6 + 2))


def _scale(ws, rng):
    ws.conditional_formatting.add(rng, ColorScaleRule(start_type='num', start_value=0, start_color='F8696B',
                                                      mid_type='num', mid_value=0.6, mid_color='FFEB84',
                                                      end_type='num', end_value=1, end_color='63BE7B'))


def write_table(ws, df: pd.DataFrame, row: int, col: int = 1, *, filt: bool = False, pct=(), sim=None,
                fills: dict | None = None, widths: bool = True) -> int:
    """df를 (row, col)부터 쓰고 다음 빈 행 번호를 돌려준다. sim: 색 척도를 줄 열(기본: 이름에 유사도/일치도/일치율)."""
    if df is None or df.empty:
        ws.cell(row=row, column=col, value='해당 없음').font = F_NOTE
        return row + 2
    cols = list(df.columns)
    for j, name in enumerate(cols):
        c = ws.cell(row=row, column=col + j, value=name)
        c.font, c.fill, c.alignment, c.border = F_HEAD, FILL_HEAD, CENTER, BORDER
    for i, rec in enumerate(df.itertuples(index=False), start=1):
        for j, v in enumerate(rec):
            if isinstance(v, float) and pd.isna(v):
                v = None
            c = ws.cell(row=row + i, column=col + j, value=v)
            c.font, c.border = F_BASE, BORDER
            name = cols[j]
            c.alignment = WRAP if name in LONG_TEXT or (isinstance(v, str) and len(v) > 30) else Alignment(vertical='top')
            if name in pct:
                c.number_format = '0%'
            elif isinstance(v, float):
                c.number_format = '0.00'
            if fills and name in fills and v in fills[name]:
                c.fill = PatternFill('solid', fgColor=fills[name][v])
    last = row + len(df)
    sim_cols = sim if sim is not None else [n for n in cols if any(w in n for w in SIM_WORDS)]
    for name in sim_cols:
        if name in cols:
            letter = get_column_letter(col + cols.index(name))
            _scale(ws, f'{letter}{row + 1}:{letter}{last}')
    if widths:
        for j, name in enumerate(cols):
            letter = get_column_letter(col + j)
            ws.column_dimensions[letter].width = max(ws.column_dimensions[letter].width or 0, _width(name, df[name]))
    if filt:
        ws.auto_filter.ref = f'{get_column_letter(col)}{row}:{get_column_letter(col + len(cols) - 1)}{last}'
    return last + 2


def _sheet(wb, title, note=None):
    ws = wb.create_sheet(title)
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value=title).font = F_SUB
    if note:
        ws.cell(row=2, column=1, value=note).font = F_NOTE
    return ws


def build_workbook(df: pd.DataFrame, res: dict, info: dict, path: str) -> str:
    wb = Workbook()
    _summary_sheet(wb.active, df, res, info)

    # 스킬 매트릭스 — 직무별 블록
    ws = _sheet(wb, '스킬 매트릭스', '직무별로 부서마다 어떤 스킬을 갖고 있는지(●). 구분: 공통=직무 내 모든 부서, 일부=2개 이상 부서, '
                                 '고유=한 부서에만. 같은 스킬 판단은 스킬명 기준(공백·대소문자 무시).')
    r = 4
    for b in res['matrix']:
        ws.cell(row=r, column=1, value=f"■ {b['job']}  (부서 {len(b['depts'])} · 스킬 {len(b['table'])})").font = F_SUB
        r = write_table(ws, b['table'], r + 1, fills={'구분': KIND_FILL}, widths=True)
        for row in ws.iter_rows(min_row=r - 1 - len(b['table']), max_row=r - 2, min_col=5, max_col=4 + len(b['depts'])):
            for c in row:
                c.alignment = CENTER
    ws.freeze_panes = 'B4'

    ws = _sheet(wb, '같은 스킬 비교', '같은 직무에서 2개 이상 부서가 가진 스킬 — 부서끼리 정의·Level 1~3 문장이 얼마나 같은지(부서 쌍 평균, 1=같음). '
                                   '가장 다르게 쓰는 스킬부터. 실제 문장은 "같은 스킬 상세" 시트.')
    write_table(ws, res['same_summary'], 4, filt=True, fills={'판정': JUDGE_FILL})
    ws.freeze_panes = 'C5'

    ws = _sheet(wb, '같은 스킬 상세', '위 스킬들의 부서별 정의·Level 1~3 원문을 나란히 — 같은 스킬끼리 묶여 있습니다.')
    detail = res['same_detail']
    end = write_table(ws, detail, 4, filt=True, fills={'판정': JUDGE_FILL})
    if not detail.empty:
        shade, prev = False, None
        for i, (j, n) in enumerate(zip(detail['직무'], detail['스킬명']), start=5):
            if (j, n) != prev:
                shade, prev = not shade, (j, n)
            if shade:
                for c in ws[i]:
                    if c.column != 3 and c.value is not None and c.fill.fgColor.rgb in (None, '00000000'):
                        c.fill = FILL_ALT
    ws.freeze_panes = 'F5'
    del end

    ws = _sheet(wb, '표준화 후보', '스킬명은 다른데 내용(스킬명+정의)이 거의 같거나, 이름 표기가 비슷한 스킬 쌍 — 스킬명 통일 검토 대상. '
                                 '내용 유사도 높은 순.')
    write_table(ws, res['standardize'], 4, filt=True)
    ws.freeze_panes = 'C5'

    ws = _sheet(wb, '부서 간 유사도', '같은 직무 안 부서끼리 스킬 구성이 얼마나 겹치는지(1=완전히 같음). 왼쪽: 스킬명이 같은 비율(Jaccard), '
                                    '오른쪽: 이름이 달라도 내용이 같은 스킬까지 같은 것으로 본 비율.')
    r = 4
    for b in res['dept']:
        ws.cell(row=r, column=1, value=f"■ {b['job']}  (평균 스킬명 일치율 {b['avg_jaccard'] or 0:.0%} · "
                                       f"내용 일치율 {b['avg_soft'] or 0:.0%})").font = F_SUB
        n = len(b['depts'])
        ws.cell(row=r + 1, column=1, value='스킬명 일치율').font = F_BOLD
        ws.cell(row=r + 1, column=n + 3, value='내용 기준 일치율').font = F_BOLD
        for k, (mat, c0) in enumerate(((b['jaccard'], 1), (b['soft'], n + 3))):
            frame = mat.reset_index().rename(columns={'index': '부서'})
            write_table(ws, frame, r + 2, c0, pct=set(b['depts']), sim=list(b['depts']), widths=False)
        r += n + 5
    for j in range(1, 40):
        ws.column_dimensions[get_column_letter(j)].width = 16

    ws = _sheet(wb, '구성요소 비율', '부서·직무별 스킬구성요소(도구/스킬/언어/지식) 개수와 비율. 한 스킬에 여러 요소가 있으면 각각 셉니다.')
    comp = res['components']
    write_table(ws, comp, 4, filt=True, pct={f'{c} %' for c in ('도구', '스킬', '언어', '지식')},
                sim=[f'{c} %' for c in ('도구', '스킬', '언어', '지식')])
    for i, d in enumerate(comp['부서'] if not comp.empty else [], start=5):
        if d == '(직무 전체)':
            for c in ws[i]:
                if c.value is not None:
                    c.font = F_BOLD
                    c.fill = FILL_BLOCK
    ws.freeze_panes = 'C5'

    ws = _sheet(wb, '레벨 점검', 'Level 1~3 빈칸, 이웃 단계 문구가 거의 같아 구분이 안 되는 경우, 단계가 올라가는데 수준 표현이 내려가는 경우'
                              '(기초: 이해·보조 / 적용: 수행·활용 / 주도: 설계·리드·지도 — 단어 기반 참고용).')
    write_table(ws, res['level'], 4, filt=True)
    ws.freeze_panes = 'D5'

    ws = _sheet(wb, '데이터 점검', '원본 정리가 필요한 행 — 필수값 누락, 같은 직무·부서 중복, 스킬명 표기 차이, 구성요소 4종 외 값.')
    write_table(ws, res['data'], 4, filt=True)
    ws.freeze_panes = 'A5'

    ws = _sheet(wb, '직무 간 공유 스킬', '2개 이상 직무에 나오는 스킬(스킬명 기준) — 공통 교육·직무 전환 시 참고.')
    write_table(ws, res['cross_job'], 4, filt=True)
    ws.freeze_panes = 'B5'

    ws = _sheet(wb, '정규화 데이터', f"분석에 쓴 원본(헤더 정규화). 원본: {info.get('source', '')}")
    cols = ['row_no', 'job_function', 'department', 'skill_name', 'skill_key', 'definition_skill',
            'skill_level_1', 'skill_level_2', 'skill_level_3', 'components_skill']
    write_table(ws, df[cols], 4, filt=True, sim=[])
    ws.freeze_panes = 'E5'

    wb.save(path)
    return path


def _summary_sheet(ws, df, res, info):
    ws.title = '요약'
    ws.sheet_view.showGridLines = False
    ws.cell(row=1, column=1, value='직무 스킬 분석').font = F_TITLE
    mode = '의미 유사도(사내 임베딩) + 문구 일치도' if res['mode'] == 'sem' else '문구 일치도(임베딩 미사용)'
    facts = [
        ('원본', info.get('source', '')),
        ('분석 일시', info.get('run_at', '')),
        ('행 수', f"{len(df)}행 (빈 스킬명 제외 {info.get('dropped_no_name', 0)}행, 병합 셀 보정 {info.get('filled_cells', 0)}칸)"),
        ('직무 / 부서 / 스킬', f"{df['job_function'].nunique()}개 / {df['department'].nunique()}개 / "
                            f"{df['skill_key'].nunique()}개(스킬명 기준)"),
        ('유사도 방식', mode + (f" — 임베딩 오류: {res['sem_error'][:120]}" if res.get('sem_error') else '')),
    ]
    r = 3
    for k, v in facts:
        ws.cell(row=r, column=1, value=k).font = F_BOLD
        ws.cell(row=r, column=2, value=v).font = F_BASE
        r += 1
    r += 1
    ws.cell(row=r, column=1, value='직무별 요약').font = F_SUB
    r = write_table(ws, res['job_summary'], r + 1, pct={'부서 간 스킬명 일치율', '부서 간 내용 일치율'})

    counts = [
        ('같은 스킬 중 "내용 다름"', int((res['same_summary']['판정'] == '내용 다름').sum()) if not res['same_summary'].empty else 0,
         '같은 스킬 비교'),
        ('표준화 후보(스킬 쌍)', len(res['standardize']), '표준화 후보'),
        ('레벨 점검 대상(행)', len(res['level']), '레벨 점검'),
        ('데이터 점검 대상', len(res['data']), '데이터 점검'),
        ('직무 간 공유 스킬', len(res['cross_job']), '직무 간 공유 스킬'),
    ]
    ws.cell(row=r, column=1, value='살펴볼 항목').font = F_SUB
    r = write_table(ws, pd.DataFrame(counts, columns=['항목', '건수', '시트']), r + 1, widths=False)

    ws.cell(row=r, column=1, value='판정 기준').font = F_SUB
    if res['mode'] == 'sem':
        crit = '의미 유사도 0.85 이상 = 같은 내용, 0.70~0.85 = 일부 다름, 0.70 미만 = 내용 다름 (정의·Level 1~3 평균)'
    else:
        crit = '문구 일치도 0.80 이상 = 같은 내용, 0.50~0.80 = 일부 다름, 0.50 미만 = 내용 다름 (정의·Level 1~3 평균)'
    notes = [crit,
             '문구 일치도: 글자 조각이 얼마나 겹치는지(표현 차이에 민감). 의미 유사도: 표현이 달라도 뜻이 같으면 높게 나옴.',
             '같은 스킬 판단은 스킬명 기준(공백·대소문자·구두점 무시). 이름이 다르지만 같은 스킬로 보이는 경우는 "표준화 후보".',
             '레벨 점검의 수준 표현은 단어 기반 참고용입니다 — 실제 판단은 원문을 확인하세요.']
    for n in notes:
        r += 1
        ws.cell(row=r, column=1, value='• ' + n).font = F_BASE
    ws.column_dimensions['A'].width = 26
    ws.column_dimensions['B'].width = 16
