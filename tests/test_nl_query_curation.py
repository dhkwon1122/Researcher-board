import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'pipeline'))
sys.path.insert(0, ROOT)

import pytest  # noqa: E402
from openpyxl import load_workbook  # noqa: E402

import llm_client  # noqa: E402
from services import ai_search_lab as lab  # noqa: E402
from services import data_store, open_data_query, query_settings  # noqa: E402
from services import nl_query_curation as cur  # noqa: E402
from services import web_pipeline_runner as wpr  # noqa: E402


@pytest.fixture()
def env(tmp_path, monkeypatch):
    monkeypatch.setattr(data_store, 'DATA_DIR', str(tmp_path / 'data'))
    monkeypatch.setattr(query_settings, '_RULES_PATH', str(tmp_path / 'data' / 'rules.txt'))
    monkeypatch.setattr(wpr, 'WEB_UPDATES_DIR', str(tmp_path / 'web'))
    return tmp_path


def _entry(**kw):
    base = {'row': 2, 'question': 'Q1', 'category': '복합 조건', 'status': '실패', 'kind': '규칙', 'approach': '',
            'sql': '', 'expected_ids': [], 'rule_text': '', 'sys_sql': '', 'intent': '', 'judge_reason': '',
            'judge_issue': ''}
    base.update(kw)
    return base


def test_rule_block_preserves_manual_rules_and_roundtrips(env):
    query_settings.write_rules('수기 규칙 A')
    res = cur.apply_entries(cur.validate([_entry(rule_text='상위평가는 가/나 등급이다')], run_sql=False), 'adm')
    assert res['rules'] == 1
    text = query_settings.read_rules()
    assert '수기 규칙 A' in text and '- 상위평가는 가/나 등급이다' in text and cur._BLOCK_START in text
    # 같은 규칙 재반영은 중복으로 건너뜀
    res2 = cur.apply_entries(cur.validate([_entry(rule_text='상위평가는 가/나 등급이다')], run_sql=False), 'adm')
    assert res2['rules'] == 0 and res2['skipped_duplicate_rules'] == 1
    rid = cur.list_rules()[0]['id']
    cur.set_active('rule', [rid], False)
    assert '상위평가' not in query_settings.read_rules() and '수기 규칙 A' in query_settings.read_rules()
    cur.set_active('rule', [rid], True)
    cur.delete('rule', [rid])
    assert query_settings.read_rules().strip() == '수기 규칙 A' and cur.list_rules() == []


def test_rule_limit_rejects_without_partial_apply(env):
    query_settings.write_rules('x' * (query_settings._MAX_LEN - 50))
    with pytest.raises(ValueError):
        cur.apply_entries(cur.validate([_entry(rule_text='y' * 200)], run_sql=False), 'adm')
    assert cur.list_rules() == []


def test_validation_rules(env):
    es = cur.validate([
        _entry(kind='규칙', rule_text=''),
        _entry(kind='예시', approach='', sql=''),
        _entry(kind='예시', sql='DROP TABLE researchers'),
        _entry(kind='예시', sql='SELECT researcher_id FROM researchers'),
        _entry(kind='코드', approach=''),
        _entry(kind='모름', approach='x'),
    ], run_sql=False)
    assert [e['ok'] for e in es] == [False, False, False, True, False, False]


def test_example_hint_exact_match_works_without_embeddings(env, monkeypatch):
    import researcher_fit as fit
    monkeypatch.setattr(fit, 'cached_embed', lambda texts: (_ for _ in ()).throw(RuntimeError('no embed server')))
    cur.apply_entries(cur.validate([_entry(kind='예시', question='박사 연구원 몇 명', approach='학력 박사 기준 distinct',
                                           sql='SELECT COUNT(DISTINCT researcher_id) FROM education')], run_sql=False), 'adm')
    hint = cur.examples_hint_for('박사  연구원 몇 명')
    assert '올바른 접근: 학력 박사 기준 distinct' in hint and 'SELECT COUNT' in hint
    assert '분류: open_data_query' in cur.examples_hint_for('박사 연구원 몇 명', for_router=True)
    assert cur.examples_hint_for('전혀 다른 질문') == ''
    eid = cur.list_examples()[0]['id']
    cur.set_active('example', [eid], False)
    assert cur.examples_hint_for('박사 연구원 몇 명') == ''


def test_hint_reaches_sql_generation_prompt(env, monkeypatch):
    cur.apply_entries(cur.validate([_entry(kind='예시', question='질문Z', approach='접근Z')], run_sql=False), 'adm')
    seen = {}
    monkeypatch.setattr(llm_client, 'call_llm', lambda prompt, system, **k: seen.setdefault('system', system) and '')
    open_data_query._generate_sql('질문Z', 'schema', None)
    assert '접근Z' in seen['system'] and '관리자가 검증한 예시' in seen['system']


def test_excel_roundtrip_code_request_and_golden(env):
    run = {'label': 'r', 'started': 't', 'finished': 't', 'status': 'done', 'results': [{
        'category': '복합 조건', 'question': 'Q1', 'status': '실패', 'intent': 'open_data_query', 'row_count': 1,
        'total_rows': 1, 'seconds': 1.0, 'flags': [], 'golden': None, 'judge': {'score': 2, 'issue_type': '조건 누락', 'reason': '박사 조건 빠짐'},
        'sql': 'SELECT 1', 'sample_rows': [], 'answer': '', 'note': ''}]}
    wb = load_workbook(io.BytesIO(lab.build_run_workbook(run)))
    ws = wb['결과']
    headers = [c.value for c in ws[1]]
    assert headers[-5:] == cur.CURATION_HEADERS and '작성 방법' in wb.sheetnames
    row = {h: i + 1 for i, h in enumerate(headers)}
    ws.cell(row=2, column=row[cur.H_KIND], value='코드')
    ws.cell(row=2, column=row[cur.H_APPROACH], value='학위 조건을 education 조인으로 처리해야 함')
    ws.cell(row=2, column=row[cur.H_IDS], value='1, 00000002\n3.0')
    buf = io.BytesIO()
    wb.save(buf)
    entries, notes = cur.parse_upload(buf.getvalue())
    assert len(entries) == 1 and entries[0]['expected_ids'] == ['00000001', '00000002', '00000003']
    assert entries[0]['judge_reason'] == '박사 조건 빠짐'
    res = cur.apply_entries(cur.validate(entries, run_sql=False), 'adm', 'run1')
    assert res['code'] == 1 and res['golden'] == 1
    assert '코드 수정 요청서' in res['code_request'] and 'education 조인' in res['code_request'] and 'SELECT 1' in res['code_request']
    draft = lab.load_draft()
    assert draft[0]['question'] == 'Q1' and draft[0]['expected']['ids'] == ['00000001', '00000002', '00000003']


def test_upload_requires_expected_columns():
    wb_bytes = io.BytesIO()
    from openpyxl import Workbook
    wb = Workbook()
    wb.active.append(['아무거나'])
    wb.save(wb_bytes)
    entries, notes = cur.parse_upload(wb_bytes.getvalue())
    assert entries == [] and '질문' in notes[0]
