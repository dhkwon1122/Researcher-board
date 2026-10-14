"""엑셀 금지 제어문자 제거(2026-10-14 — 특허 제목의 제어문자로 연구원 명단 엑셀 다운로드 전체가 실패)."""
import io

from openpyxl import Workbook, load_workbook

from services.excel_text import excel_safe


def test_excel_safe_strips_control_chars_keeps_newline_tab():
    raw = '2026-01 : P-G\x0b 특허\x1f 데이터\n둘째 줄\t탭'
    assert excel_safe(raw) == '2026-01 : P-G 특허 데이터\n둘째 줄\t탭'
    assert excel_safe(3) == 3 and excel_safe(None) is None


def test_cleaned_value_can_be_saved():
    wb = Workbook()
    wb.active.cell(row=1, column=1, value=excel_safe('특허\x0b제목'))
    buf = io.BytesIO()
    wb.save(buf)
    assert load_workbook(io.BytesIO(buf.getvalue())).active['A1'].value == '특허제목'
