"""JOB Market 심화 지표(services/job_market_insights.py, 2026-10-13) — 합성 데이터로 검증."""
import os
import sys
import numpy as np
import pandas as pd

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, ROOT)
sys.path.insert(0, os.path.join(ROOT, 'pipeline'))

from services import job_market_insights as jmi  # noqa: E402

# 단어별 고정 축: 같은 단어를 공유하면 유사도가 높다.
_VOCAB = ['광학', '소재', '배터리', 'AI', '설계']


def _embed(texts):
    out = []
    for t in texts:
        v = np.array([1.0 if w in t else 0.0 for w in _VOCAB]) + 1e-3
        out.append(v / np.linalg.norm(v))
    return np.array(out)


def _ctx(monkeypatch, profiles, gap, edges=None, researchers=None, contrib=None, extra=None):
    from services import data_store as ds
    tables = {
        **(extra or {}),
        'researchers': researchers if researchers is not None else pd.DataFrame(
            {'researcher_id': ['00000001', '00000002', '00000003', '00000004'],
             'name': ['갑', '을', '병', '정'], 'org_code': ['P', 'X', 'X', 'Y'], 'position': ['CL3'] * 4}),
        'collaboration_edges': edges if edges is not None else pd.DataFrame(
            [{'researcher_a': '00000001', 'researcher_b': '00000002', 'recent_count': 2, 'total_count': 2}]),
        'researcher_contribution_metrics': contrib if contrib is not None else pd.DataFrame(
            {'researcher_id': ['00000002'], 'contribution_type': ['주도형']}),
    }
    monkeypatch.setattr(ds, 'read_processed', lambda name, **k: tables.get(name, pd.DataFrame()))
    monkeypatch.setattr(ds, 'read_project_competency_gap', lambda: gap)
    monkeypatch.setattr(jmi, '_tenure_map', lambda df: {'00000001': 'Senior', '00000002': 'Junior', '00000003': 'Junior'})
    return jmi.Context(profiles, embed_fn=_embed)


GAP = [{'project_name': '[연구]X과제', 'competencies': [
    {'competency': '광학 설계', 'description': '빛 설계',
     'members': [{'researcher_id': '00000002', 'score': 0.4}, {'researcher_id': '00000003', 'score': 0.3}],
     'outsiders': [{'researcher_id': '00000001', 'score': 0.9, 'reason': 'r'},
                   {'researcher_id': '00000004', 'score': 0.8}]},
    {'competency': '배터리 소재', 'members': [{'researcher_id': '00000002', 'score': 0.9}], 'outsiders': []},
]}]
PROFILES = {'00000001': {'strength_fields': ['광학 설계'], 'current_focus_fields': ['광학'], 'past_fields': ['배터리']},
            '00000004': {'strength_fields': ['AI']}}


def test_gap_fill_and_evidence(monkeypatch):
    ctx = _ctx(monkeypatch, PROFILES, GAP)
    g = jmi.gap_fill(ctx, '00000001', 'X과제')   # 꼬리표 무시 매칭
    assert g['weak_count'] == 1 and g['total'] == 2
    assert [f['competency'] for f in g['filled']] == ['광학 설계']
    assert g['filled'][0]['evidence'] == '광학 설계' and g['score'] > 0.5
    assert jmi.gap_fill(ctx, '00000004', 'X과제')['filled'] == []
    assert jmi.gap_fill(ctx, '00000001', '없는과제') is None


def test_focus_collab_composition_and_ranking(monkeypatch):
    ctx = _ctx(monkeypatch, PROFILES, GAP)
    proj = _embed(['광학'])[0]
    assert jmi.focus_match(ctx, '00000001', proj)['type'] == 'current'
    assert jmi.focus_match(ctx, '00000001', _embed(['배터리'])[0])['type'] == 'past'
    assert jmi.collaborators(ctx, '00000001', ['00000002', '00000003']) == [('00000002', '을')]
    comp = jmi.composition(ctx, '00000001', ['00000002', '00000003'])
    assert comp['seniors'] == 0 and comp['complement'] == '시니어 보완' and comp['leads'] == 1
    scored = [{'project_name': '다른과제', 'dep_name': 'D', 'combined': 0.70, 'a_embedding': None},
              {'project_name': 'X과제', 'dep_name': 'D', 'combined': 0.65, 'a_embedding': proj}]
    out = jmi.enrich_candidates(ctx, '00000001', scored, {'X과제': ['00000002', '00000003']})
    assert out[0]['project_name'] == 'X과제'           # 갭 충족 가산으로 역전
    text = jmi.signals_text(out[0])
    assert '광학 설계' in text and '을' not in text      # 프롬프트에 이름(PII) 없음
    assert jmi.public_fields(out[0])['gap']['filled']


def test_growth_path(monkeypatch):
    ctx = _ctx(monkeypatch, PROFILES, GAP)
    nur = pd.DataFrame([{'researcher_id': '00000002', 'category': '학위', 'subcategory': '박사', 'institution': 'KAIST',
                         'major': '재료'}])
    g = jmi.growth_path(ctx, '00000004', 'X과제', nurturing=nur)
    assert {m['competency'] for m in g['missing']} == {'광학 설계', '배터리 소재'}
    bat = [m for m in g['missing'] if m['competency'] == '배터리 소재'][0]
    assert bat['mentors'][0]['name'] == '을' and 'KAIST' in bat['mentors'][0]['nurturing'][0]


def test_plan_assignment_capacity():
    results = {'a': {'shortlist': [{'project_name': 'P1', 'adjusted': 0.9}, {'project_name': 'P2', 'adjusted': 0.5}]},
               'b': {'shortlist': [{'project_name': 'P1', 'adjusted': 0.8}, {'project_name': 'P2', 'adjusted': 0.7}]},
               'c': {'shortlist': []}}
    plan = jmi.plan_assignment(results, capacity=1)
    got = {a['researcher_id']: a['project_name'] for a in plan['assigned']}
    assert got == {'a': 'P1', 'b': 'P2'} and plan['unassigned'] == ['c']
    assert jmi.plan_assignment(results, capacity=2)['load'] == {'P1': 2}


def test_ending_impact():
    holders = pd.DataFrame([
        {'technology': 'T1', 'source': '핵심기술', 'holder_ids': '00000001;00000002'},
        {'technology': 'T2', 'source': '핵심기술', 'holder_ids': '00000001;00000003;00000004'},
        {'technology': 'T3', 'source': '강점분야', 'holder_ids': '00000001'},
    ])
    rows = jmi.ending_impact(['00000001', '00000002'], holders)
    assert [r['technology'] for r in rows] == ['T1', 'T3']
    assert rows[0]['remaining'] == 0 and rows[0]['share'] == 1.0


def test_recommend_flow_and_render(monkeypatch):
    import json as _json

    import dash
    dash.Dash(__name__, use_pages=True, pages_folder='')
    from services import job_market as jm
    ctx = _ctx(monkeypatch, PROFILES, GAP)
    monkeypatch.setattr(jm.fit, 'cached_embed', _embed)
    prompts = []

    def fake_llm(prompt, system, **k):
        prompts.append(prompt)
        return _json.dumps({'recommendations': [{'project_name': '[연구]X과제', 'reason': '광학 설계 역량을 채울 수 있음',
                                                  'confidence': '상', 'risk': '배터리 소재 경험 부족'}],
                            'closest_non_match': None,
                            'must_place': [{'project_name': '[연구]X과제', 'reason': '약한 역량 보완에 가장 적합'}]})
    monkeypatch.setattr(jm.llm_client, 'call_llm', fake_llm)
    pool = {'[연구]X과제': {'dep_name': 'D', 'a_text': '광학', 'a_embedding': _embed(['광학'])[0],
                          'b_members': [], 'member_ids': ['00000002', '00000003']},
            '다른과제': {'dep_name': 'D', 'a_text': 'AI', 'a_embedding': _embed(['AI'])[0], 'b_members': [], 'member_ids': []}}
    res = jm.recommend_for_researcher('00000001', PROFILES, pool, ctx)
    rec = res['recommendations'][0]
    assert rec['gap']['filled'][0]['competency'] == '광학 설계' and rec['collaborators'] == [['00000002', '을']]
    assert rec['confidence'] == '상' and rec['risk'] == '배터리 소재 경험 부족'
    assert rec['coverage'] and rec['coverage']['total'] >= 2
    assert res['shortlist'][0]['project_name'] == '[연구]X과제' and res['growth'] is None
    assert 'C) 참고 지표' in prompts[0] and '00000002' not in prompts[0]
    _json.dumps(res['shortlist']), _json.dumps(rec)   # 이력 저장 가능(직렬화)

    # 추천이 없으면 육성 경로
    monkeypatch.setattr(jm.llm_client, 'call_llm', lambda *a, **k: _json.dumps(
        {'recommendations': [], 'closest_non_match': {'project_name': '[연구]X과제', 'reason': '전문 분야가 다름'},
         'must_place': []}))
    res2 = jm.recommend_for_researcher('00000004', PROFILES, pool, ctx)
    assert res2['growth'] and res2['growth']['missing']

    import pages.job_market as page
    monkeypatch.setattr(page, '_eval_cells', lambda roster: None)
    report = {'mode': 'project', 'project_names': ['P'], 'run_at': 'now', 'candidates_considered': 2,
              'roster': [{'researcher_id': '00000001', 'name': '갑', 'department': 'd', 'org_code': 'P',
                          'position_year': 'CL3', 'degree_major': '-', 'age': 40},
                         {'researcher_id': '00000004', 'name': '정', 'department': 'd', 'org_code': 'Y',
                          'position_year': 'CL3', 'degree_major': '-', 'age': 30}],
              'results': {'00000001': res, '00000004': res2},
              'impact': [{'technology': 'T1', 'source': 's', 'holders': 2, 'ending_holders': 2, 'remaining': 0, 'share': 1.0}]}
    html_text = str(page._render_result(report))
    for needle in ('갭 충족 1/1', '기존 협업자 1명', '일괄 배치 시뮬레이션', '종료 영향 분석', '육성 경로',
                   '확신도 상', '부족한 점: 배터리 소재 경험 부족', '요구사항 충족'):
        assert needle in html_text, needle
    import io

    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(jm.build_result_workbook(report))).active
    cell = str(ws.cell(row=2, column=2).value)
    assert '갭 충족: 광학 설계' in cell and '확신도 상' in cell and '부족한 점: 배터리 소재 경험 부족' in cell



def test_project_requirements_split():
    entry = {'core_tech': '광학 설계 기반 센서. 소재 합성 공정 개발', 'deliverable': '확인 불가', 'challenge': '',
             'keywords_kr': ['배터리', '광학 설계 기반 센서'], 'keywords_en': ['battery']}
    reqs = jmi.project_requirements(entry)
    assert reqs == ['광학 설계 기반 센서', '소재 합성 공정 개발', '배터리']   # 문장 분리, 확인 불가 제외, 중복 제거, 한글 키워드 우선
    assert jmi.project_requirements(None) == []


def test_requirement_coverage_and_evidence(monkeypatch):
    from datetime import date
    y = date.today().year
    pubs = pd.DataFrame([{'researcher_id': '00000001', 'title': '광학 설계 최적화 연구', 'pub_year': str(y - 1), 'pub_date': ''},
                         {'researcher_id': '00000001', 'title': 'AI 모델', 'pub_year': str(y - 9), 'pub_date': ''}])
    pats = pd.DataFrame([{'researcher_id': '00000001', 'title': '', 'title_ko': '배터리 소재 특허', 'application_date': f'{y}0101'}])
    tasks = pd.DataFrame([{'researcher_id': '00000001', 'task_name': '광학 과제', 'the_task_name': '', 'end_date': ''}])
    profiles = {'00000001': {'strength_fields': ['광학 설계'], 'key_responsibilities': ['소재 합성']}}
    ctx = _ctx(monkeypatch, profiles, GAP, extra={'publications': pubs, 'patents': pats, 'tasks': tasks})
    cov = jmi.requirement_coverage(ctx, '00000001', 'X과제', ['광학 설계 센서', '소재 공정', 'AI 반도체'])
    # 요구 = 기본 3개 + 필요 역량 2개(광학 설계, 배터리 소재)
    assert cov['total'] == 5
    covered = {c['requirement']: c['item'] for c in cov['covered']}
    assert {'광학 설계 센서', '소재 공정', '광학 설계'} <= set(covered)
    assert covered['소재 공정'] == '소재 합성'                         # 역할·책임도 충족 근거
    assert 'AI 반도체' in cov['uncovered'] and cov['ratio'] == round(len(covered) / 5, 3)
    ev = jmi.evidence_match(ctx, '00000001', _embed(['광학 설계'])[0])
    kinds = {(x['kind'], x['title']) for x in ev['items']}
    assert ('논문', '광학 설계 최적화 연구') in kinds and ('과제', '광학 과제') in kinds
    assert all(x['title'] != 'AI 모델' for x in ev['items'])           # 5년 이전 실적 제외
    assert jmi.evidence_match(ctx, '00000001', None) is None
