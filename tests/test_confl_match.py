import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'pipeline'))
sys.path.insert(0, ROOT)

from openpyxl import Workbook, load_workbook  # noqa: E402

from services import confl_match  # noqa: E402

PAGES = [
    {'depth': 0, 'id': '1', 'title': "월별 '26.9", 'parent_id': '', 'parent_title': ''},
    {'depth': 1, 'id': '10', 'title': 'AI팀', 'parent_id': '1', 'parent_title': "월별 '26.9"},
    {'depth': 2, 'id': '3862782334', 'title': "2D Material|'26.9월", 'parent_id': '10', 'parent_title': 'AI팀'},
    {'depth': 2, 'id': '3000000111', 'title': "SLAM 개발|'26.9월", 'parent_id': '10', 'parent_title': 'AI팀'},
    {'depth': 2, 'id': '3000000222', 'title': "SLAM 개발|'26.9월", 'parent_id': '11', 'parent_title': '로봇팀'},
    {'depth': 2, 'id': '3000000333', 'title': "신규과제|'26.9월", 'parent_id': '11', 'parent_title': '로봇팀'},
    {'depth': 3, 'id': '3000000444', 'title': '월간보고서', 'parent_id': '3862782334', 'parent_title': '2D Material'},
]


def _source(path, extra_cols=0):
    wb = Workbook()
    ws = wb.active
    ws.append(['소속', '과제명', '컨플 주소'] + [f'기타{i}' for i in range(extra_cols)])
    ws.append(['AI팀', '[연구]2D Material', '999'])         # 기존 값 → 덮어씀
    ws.append(['로봇팀', 'SLAM 개발', ''])
    ws.append(['AI팀', 'SLAM 개발', ''])
    ws.append(['AI팀', '완전히 다른 과제', '555'])           # 실패 → 기존 값 유지
    wb.save(path)


def test_normalize_title():
    assert confl_match.normalize_title("[연구]2D Material") == confl_match.normalize_title("2D Material|'26.9월")
    assert confl_match.similarity('[연구]2D Material', "2D Material|'26.9월") == 1.0
    assert confl_match.similarity('완전히 다른 과제', 'SLAM 개발') < confl_match.MIN_SIMILARITY


def test_apply_matches_by_title_and_department(tmp_path, monkeypatch):
    src = tmp_path / '과제별컨플.xlsx'
    _source(src)
    monkeypatch.setattr(confl_match, 'find_source_workbook', lambda: str(src))
    data, summary, source = confl_match.build_updated_workbook(PAGES)
    ws = load_workbook(io.BytesIO(data)).worksheets[0]
    rows = {r[1].value: r for r in ws.iter_rows(min_row=2)}
    assert summary == {'success': 3, 'fail': 1, 'unmatched_pages': 1, 'candidates': 4, 'result_column': 'L'}
    assert ws['L1'].value == '추출 결과' and ws['M1'].value == '매칭된 제목'
    r2 = [r for r in ws.iter_rows(min_row=2) if r[0].value == 'AI팀' and '2D' in r[1].value][0]
    assert r2[2].value == '3862782334' and r2[11].value == '성공'      # 덮어쓰기, 텍스트
    slam = {r[0].value: r for r in ws.iter_rows(min_row=2) if r[1].value == 'SLAM 개발'}
    assert slam['AI팀'][2].value == '3000000111' and slam['로봇팀'][2].value == '3000000222'   # 소속으로 구분
    bad = [r for r in ws.iter_rows(min_row=2) if r[1].value == '완전히 다른 과제'][0]
    assert bad[2].value == '555' and bad[11].value == '실패'             # 기존 값 유지
    leftover = load_workbook(io.BytesIO(data))['미매칭 페이지']
    assert [r[0].value for r in leftover.iter_rows(min_row=2)] == ["신규과제|'26.9월"]


def test_result_column_goes_after_wide_files_and_is_reused(tmp_path, monkeypatch):
    src = tmp_path / '과제별컨플.xlsx'
    _source(src, extra_cols=12)           # 15열 → 결과는 P열(16)
    monkeypatch.setattr(confl_match, 'find_source_workbook', lambda: str(src))
    data, summary, _ = confl_match.build_updated_workbook(PAGES)
    assert summary['result_column'] == 'P'
    src2 = tmp_path / 'again.xlsx'
    src2.write_bytes(data)
    monkeypatch.setattr(confl_match, 'find_source_workbook', lambda: str(src2))
    _, summary2, _ = confl_match.build_updated_workbook(PAGES)
    assert summary2['result_column'] == 'P'   # 이미 있는 "추출 결과" 열을 다시 쓴다


def test_missing_inputs_raise_value_error(monkeypatch):
    import pytest
    with pytest.raises(ValueError):
        confl_match.build_updated_workbook([])
