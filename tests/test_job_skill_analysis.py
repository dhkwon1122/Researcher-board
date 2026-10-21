"""직무 스킬 분석(로컬, 2026-10-21) — 정규화·분석·엑셀 리포트·실행 스크립트."""
import hashlib
import os
import re
import sys

import numpy as np
import pandas as pd
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'pipeline'))
sys.path.insert(0, ROOT)
import job_skill_analysis as jsa  # noqa: E402
import process_job_skill_set as prep  # noqa: E402

HEAD = ['스킬명', '부서', '직무', '스킬 정의', 'Level 1', 'Level2', 'LEVEL3', '스킬 구성요소']
ROWS = [
    ['Python 프로그래밍', 'A부서', '회로설계', '파이썬으로 자동화 코드를 작성하는 능력', '기본 문법을 이해한다', '자동화 코드를 구현한다', '코드 구조를 설계한다', '언어'],
    ['회로 시뮬레이션', '', '', 'SPICE로 회로 동작을 검증하는 능력', '지도하에 수행한다', '결과를 분석한다', '검증 전략을 주도한다', '도구, 지식'],
    ['파이썬 코딩', 'B부서', '회로설계', '파이썬으로 자동화 코드를 작성하는 역량', '문법을 이해한다', '코드를 작성한다', '품질 기준을 정립한다', '언어'],
    ['회로 시뮬레이션', '', '', '아날로그 회로 성능을 예측하는 능력', '검증 전략을 주도한다', '시뮬레이션을 수행한다', '시뮬레이션을 수행한다', '기술'],
    ['회로시뮬레이션', '', '', '아날로그 회로 성능을 예측하는 능력', '', '', '', '도구'],
    ['데이터 분석', 'A부서', '소재개발', '실험 데이터를 분석하는 능력', '이해한다', '활용한다', '주도한다', '스킬'],
    ['데이터 분석', 'A부서', '회로설계', '측정 데이터를 분석하는 능력', '이해한다', '활용한다', '주도한다', '스킬'],
]


def _raw():
    return pd.DataFrame(ROWS, columns=HEAD)


def _fake_embed(texts):
    """동의어를 같은 토큰으로 바꾼 단어 주머니 벡터 — 의미 유사도 경로 시험용."""
    syn = {'파이썬': 'python', '코딩': '프로그래밍', '역량': '능력'}
    out = []
    for t in texts:
        s = t.lower()
        for a, b in syn.items():
            s = s.replace(a, b)
        v = np.zeros(128, dtype=np.float32)
        for tok in re.findall(r'[0-9a-z가-힣]+', s):
            v[int(hashlib.md5(tok.encode()).hexdigest(), 16) % 128] += 1
        out.append(v)
    return out


def test_normalize_headers_filldown_and_key():
    df, info = prep.normalize(_raw())
    assert list(df.columns[:8]) == prep.COLUMNS
    assert info['filled_cells'] == 6                                  # 병합 셀(빈 부서·직무) 보정
    assert df.loc[1, 'department'] == 'A부서' and df.loc[3, 'job_function'] == '회로설계'
    assert prep.skill_key('회로 시뮬레이션') == prep.skill_key('회로시뮬레이션')
    with pytest.raises(ValueError, match='현재 파일 헤더'):
        prep.normalize(_raw().drop(columns=['스킬 정의']))


def test_analysis_text_mode():
    df, _ = prep.normalize(_raw())
    res = jsa.analyze(df, embed_fn=None, log=lambda *a: None)
    assert res['mode'] == 'char'
    circuit = next(b for b in res['matrix'] if b['job'] == '회로설계')
    kinds = dict(zip(circuit['table']['스킬명'], circuit['table']['구분']))
    assert kinds['회로 시뮬레이션'] == '공통' and kinds['Python 프로그래밍'] == '고유'
    same = res['same_summary']
    assert set(same['스킬명']) == {'회로 시뮬레이션'} and same.iloc[0]['부서 수'] == 2
    assert not res['same_detail'].empty
    level = res['level']
    assert any('역전' in x for x in level['점검 결과'])                  # 주도 → 적용 → 적용
    assert any('L2·L3 문구 거의 동일' in x for x in level['점검 결과'])
    assert any('빈칸: L1, L2, L3' in x for x in level['점검 결과'])
    kinds = set(res['data']['유형'])
    assert {'중복 행', '스킬명 표기 다름', '구성요소 미분류'} <= kinds
    assert list(res['cross_job']['스킬명']) == ['데이터 분석']
    comp = res['components']
    a = comp[(comp['직무'] == '회로설계') & (comp['부서'] == 'A부서')].iloc[0]
    assert a['도구'] == 1 and a['지식'] == 1 and a['언어'] == 1 and a['스킬'] == 1
    assert not res['dept'][0]['jaccard'].empty


def test_analysis_semantic_mode_finds_synonym_skill():
    df, _ = prep.normalize(_raw())
    res = jsa.analyze(df, embed_fn=_fake_embed, log=lambda *a: None)
    assert res['mode'] == 'sem'
    std = res['standardize']
    pair = {frozenset((a, b)) for a, b in zip(std['스킬명 A'], std['스킬명 B'])}
    assert frozenset(('Python 프로그래밍', '파이썬 코딩')) in pair
    blk = next(b for b in res['dept'] if b['job'] == '회로설계')
    assert blk['soft'].loc['A부서', 'B부서'] > blk['jaccard'].loc['A부서', 'B부서']   # 이름 달라도 내용 같은 스킬 반영


def test_embedding_failure_falls_back_to_text():
    df, _ = prep.normalize(_raw())

    def broken(_texts):
        raise RuntimeError('서버 없음')
    msgs = []
    res = jsa.analyze(df, embed_fn=broken, log=msgs.append)
    assert res['mode'] == 'char' and '서버 없음' in res['sem_error']
    assert any('임베딩 서버 사용 불가' in m for m in msgs)


def test_report_and_script(tmp_path, monkeypatch):
    raw_dir, raw_csv_dir, out_dir = tmp_path / 'raw', tmp_path / 'raw_csv', tmp_path / 'processed'
    for d in (raw_dir, raw_csv_dir, out_dir):
        d.mkdir()
    monkeypatch.setattr(prep, 'RAW_DIR', str(raw_dir))
    monkeypatch.setattr(prep, 'RAW_CSV_PATH', str(raw_csv_dir / 'job_skill_set.csv'))
    monkeypatch.setattr(prep, 'OUT_DIR', str(out_dir))
    monkeypatch.setattr(prep, 'OUT_PATH', str(out_dir / 'job_skill_set.csv'))
    _raw().to_csv(raw_csv_dir / 'job_skill_set.csv', index=False, encoding='utf-8-sig')
    sys.path.insert(0, os.path.join(ROOT, 'scripts'))
    import analyze_job_skills
    assert analyze_job_skills.main(['--no-embed', '--out', str(tmp_path / 'report')]) == 0
    files = os.listdir(tmp_path / 'report')
    assert len(files) == 1 and files[0].startswith('job_skill_analysis_')
    from openpyxl import load_workbook
    wb = load_workbook(tmp_path / 'report' / files[0])
    assert wb.sheetnames == ['요약', '스킬 매트릭스', '같은 스킬 비교', '같은 스킬 상세', '표준화 후보', '부서 간 유사도',
                             '구성요소 비율', '레벨 점검', '데이터 점검', '직무 간 공유 스킬', '정규화 데이터']
    assert (out_dir / 'job_skill_set.csv').exists()
    assert analyze_job_skills.main(['--no-embed', '--job', '없는직무', '--out', str(tmp_path / 'r2')]) == 1
