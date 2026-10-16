"""엑셀 다운로드 권한 + 1회성 요청/승인(2026-10-15, 2026-10-16 개정: 기본 불허·헤더 구성 첨부·로그)."""
import io
import os
import sys

import pytest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from services import download_access as da  # noqa: E402

USER = {'user_id': 'u1', 'display_name': '홍길동', 'email': '', 'role': 'talent_dev', 'is_admin': False,
        'permissions': {}, 'eval_excluded_dep_ids': []}
ADMIN = {'user_id': 'adm', 'display_name': '관리자', 'email': '', 'role': 'talent_mng', 'is_admin': True,
         'permissions': {}, 'eval_excluded_dep_ids': []}
RB_CTX = {'ids': ['R1', 'R2'], 'columns': [{'header': '이름', 'fields': ['name'], 'join': 'paren'},
                                           {'header': '', 'fields': [], 'join': 'paren'}], 'max_items': 3}


@pytest.fixture
def paths(tmp_path, monkeypatch):
    monkeypatch.setattr(da, 'REQUESTS_PATH', str(tmp_path / 'r.json'))
    monkeypatch.setattr(da, 'LOG_PATH', str(tmp_path / 'log.jsonl'))
    monkeypatch.setattr(da, 'PREFILL_PATH', str(tmp_path / 'p.json'))


def test_role_default_denies_admin_allows(paths):
    from services import auth
    with auth.acting_as(USER):                               # 역할 기본값 = 불허
        assert not da.has_permission()
        assert not da.check('researcher_list')[0]
    with auth.acting_as({**USER, 'permissions': {'download_excel': True}}):
        assert da.check('researcher_list') == (True, '', None)
    with auth.acting_as(ADMIN):
        assert da.has_permission()


def test_request_with_spec_approve_ack_use_once(paths):
    from services import auth
    with auth.acting_as(USER):
        ok, msg, _ = da.check('report_builder')
        assert not ok and '다운로드 요청' in msg
        with pytest.raises(ValueError):
            da.create_request('report_builder', '', '')               # 사용목적 없음
        with pytest.raises(ValueError):
            da.create_request('report_builder', '', '검토', {'ids': [], 'columns': RB_CTX['columns']})  # 대상 없음
        with pytest.raises(ValueError):
            da.create_request('researcher_list', '', '검토', RB_CTX)  # 연구원 명단은 조서 작성으로
        req = da.create_request('report_builder', 'AI플랫폼', '재배치 검토', RB_CTX)
        assert req['headers'] == ['이름'] and req['target_count'] == 2   # 빈 열은 빠짐
        assert '승인 대기' in da.check('report_builder')[1]
        with pytest.raises(PermissionError):
            da.decide(req['id'], True)
    with auth.acting_as(ADMIN):
        assert da.decide(req['id'], True, '회의용')['status'] == 'approved'
        with pytest.raises(ValueError):
            da.decide(req['id'], False)                      # 이미 처리됨
    with auth.acting_as(USER):
        assert [n['id'] for n in da.notices('u1')] == [req['id']]
        ok, msg, _ = da.check('report_builder')
        assert not ok and '안내를 확인' in msg                # 책임 안내 확인 전에는 사용 불가
        da.acknowledge([req['id']], 'u1')
        assert not da.check('job_market')[0]                 # 다른 화면에는 쓸 수 없음
        ok, _msg, grant = da.check('report_builder')
        assert ok and grant['id'] == req['id'] and grant['spec']['ids'] == ['R1', 'R2']
        assert grant['spec']['max_items'] == 3
        da.complete('report_builder', grant, 2, ['이름'], 'x.xlsx')
        assert not da.check('report_builder')[0]             # 1회만
    events = [e['event'] for e in reversed(da.read_log())]
    assert events == ['request', 'approve', 'acknowledge', 'download']
    assert da.read_log()[0]['basis'] == '승인 1회' and da.read_log()[0]['requester_id'] == 'u1'


def test_fixed_headers_request(paths):
    from services import auth
    with auth.acting_as(USER):
        with pytest.raises(ValueError):
            da.create_request('expertise', '', '검토', {'ids': ['R1']}, [])
        req = da.create_request('expertise', '', '검토', {'ids': ['R1', 'R1']}, ['유사 연구원 명단', '사번', '없는 헤더'])
        assert req['spec'] == {'ids': ['R1'], 'headers': ['사번', '유사 연구원 명단']}   # 엑셀 열 순서


def test_workbooks_filter_headers():
    from openpyxl import load_workbook
    from services import job_market as jm
    result = {'roster': [{'researcher_id': 'A'}, {'researcher_id': 'B'}],
              'results': {'A': {'recommendations': [], 'must_place': []},
                          'B': {'recommendations': [], 'must_place': [{'project': 'X'}]}}}
    assert jm.result_researcher_ids(result) == ['B']
    ws = load_workbook(io.BytesIO(jm.build_result_workbook(result, ['B'], ['사번']))).active
    assert [c.value for c in ws[1]] == ['사번'] and ws['A2'].value == 'B' and ws.max_row == 2


def test_prefill_roundtrip(paths):
    da.set_prefill('u1', ['R1', 'R2', 'R1'])
    assert da.pop_prefill('u1') == ['R1', 'R2']
    assert da.pop_prefill('u1') == []
