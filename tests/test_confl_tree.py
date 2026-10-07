import io
import os
import sys

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, os.path.join(ROOT, 'pipeline'))
sys.path.insert(0, ROOT)

import confluence_client as cc  # noqa: E402
from services import confl_tree  # noqa: E402

TREE = {'100': [('101', 'A'), ('102', 'B'), ('103', 'C')], '101': [('201', 'A-1')], '201': [('301', 'A-1-i')]}
TITLES = {'9': 'Home', '100': 'Root', '101': 'A', '102': 'B', '103': 'C', '201': 'A-1', '301': 'A-1-i'}


class _Resp:
    def __init__(self, data):
        self._d, self.status_code, self.text = data, 200, ''

    def json(self):
        return self._d


class _Session:
    descendants = False   # True면 descendant/page(ancestors 포함) 지원, False면 404(→ child/page 재귀로 폴백)
    calls = 0

    def get(self, url, params=None, timeout=None):
        type(self).calls += 1
        parts = url.split('/rest/api/content/')[1].split('/')
        if len(parts) == 1:
            return _Resp({'id': parts[0], 'title': TITLES[parts[0]]})
        if parts[1] == 'descendant':
            if not self.descendants:
                r = _Resp({}); r.status_code = 404; r.text = 'not found'
                return r
            flat, parent = [], {}
            def walk(pid, chain):
                for cid, t in TREE.get(pid, []):
                    flat.append({'id': cid, 'title': t, 'ancestors': [{'id': a, 'title': TITLES[a]} for a in ['9', *chain, pid]]})
                    walk(cid, [*chain, pid])
            walk(parts[0], [])
            st, lim = (params or {}).get('start', 0), (params or {}).get('limit', 100)
            return _Resp({'results': flat[st:st + lim]})
        kids = TREE.get(parts[0], [])
        start, limit = (params or {}).get('start', 0), (params or {}).get('limit', 100)
        return _Resp({'results': [{'id': i, 'title': t} for i, t in kids[start:start + limit]]})


def _setup(monkeypatch):
    monkeypatch.setenv('CONFLUENCE_TOKEN', 'x')
    monkeypatch.setenv('CONFLUENCE_GATEWAY_BASE_URL', 'https://gw.example.net/x')
    monkeypatch.setenv('CONFLUENCE_MAX_CALLS_PER_MINUTE', '0')   # 테스트에서는 호출 간격 제한을 끈다
    _Session.calls = 0
    monkeypatch.setattr(cc, '_get_session', lambda: _Session())
    monkeypatch.setattr(cc, '_CHILD_PAGE_LIMIT', 2)   # 페이징 경로 검증


def test_crawl_all_depths_in_document_order(monkeypatch):
    _setup(monkeypatch)
    rows = cc.crawl_descendants('100.0')
    assert [r['id'] for r in rows] == ['100', '101', '201', '301', '102', '103']
    assert rows[3]['path'] == 'Root > A > A-1 > A-1-i' and rows[3]['depth'] == 3
    assert rows[2]['parent_id'] == '101'


def test_crawl_depth_limit_and_cap(monkeypatch):
    _setup(monkeypatch)
    assert [r['id'] for r in cc.crawl_descendants('100', max_depth=1)] == ['100', '101', '102', '103']
    try:
        cc.crawl_descendants('100', max_pages=2)
        assert False
    except cc.ConfluenceError:
        pass


def test_workbook_keeps_ids_as_text(monkeypatch):
    _setup(monkeypatch)
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(confl_tree.build_workbook_bytes(cc.crawl_descendants('100')))).active
    assert [c.value for c in ws[1]] == confl_tree.EXCEL_HEADERS
    assert ws['C2'].value == '100' and ws['C2'].data_type == 's'


def test_job_state_is_file_based_and_survives_across_workers(monkeypatch, tmp_path):
    import time
    from services import web_pipeline_runner as wpr
    _setup(monkeypatch)
    monkeypatch.setattr(wpr, 'WEB_UPDATES_DIR', str(tmp_path))
    ok, _ = confl_tree.start('100', None)
    assert ok
    for _ in range(100):
        if not confl_tree.is_running():
            break
        time.sleep(0.05)
    # 다른 워커 프로세스가 같은 파일을 읽는 상황 — 상태는 파일에서만 온다
    st = confl_tree.snapshot()
    assert st['status'] == 'done' and st['count'] == 6 and st['has_result']
    name, data = confl_tree.result_workbook()
    assert name.startswith('컨플_하위페이지_100_') and data[:2] == b'PK'


def test_stale_running_state_is_reported_as_error(monkeypatch, tmp_path):
    from services import web_pipeline_runner as wpr
    monkeypatch.setattr(wpr, 'WEB_UPDATES_DIR', str(tmp_path))
    confl_tree._save_state(status='running', started_at=1.0)
    st = confl_tree._read_state()
    st_path = confl_tree._state_path()
    import json
    raw = json.load(open(st_path, encoding='utf-8'))
    raw['updated_at'] = 1.0
    json.dump(raw, open(st_path, 'w', encoding='utf-8'))
    assert confl_tree.snapshot()['status'] == 'error'
    assert not confl_tree.is_running()


def test_descendant_endpoint_gives_same_tree_with_few_calls(monkeypatch):
    _setup(monkeypatch)
    monkeypatch.setattr(_Session, 'descendants', True)
    monkeypatch.setattr(cc, '_DESC_LIMIT', 4)
    rows = cc.crawl_descendants('100')
    assert [r['id'] for r in rows] == ['100', '101', '201', '301', '102', '103']
    assert rows[3]['path'] == 'Root > A > A-1 > A-1-i' and rows[3]['parent_title'] == 'A-1'
    assert _Session.calls <= 3                      # 루트 1 + 하위 목록 2(페이징), child 재귀 아님
    assert [r['id'] for r in cc.crawl_descendants('100', max_depth=1)] == ['100', '101', '102', '103']


def test_rate_limit_spacing_and_429_retry(monkeypatch):
    import time as _t
    clock = {'now': 1000.0, 'slept': []}
    monkeypatch.setattr(cc.time, 'monotonic', lambda: clock['now'])
    monkeypatch.setattr(cc.time, 'sleep', lambda s: (clock['slept'].append(s), clock.__setitem__('now', clock['now'] + s)))
    monkeypatch.setattr(cc, '_last_call_at', 0.0)
    monkeypatch.setenv('CONFLUENCE_MAX_CALLS_PER_MINUTE', '10')
    cc._throttle()                       # 첫 호출은 대기 없음
    cc._throttle()                       # 두 번째는 6초 간격
    assert clock['slept'] == [6.0]

    class Resp429:
        status_code, text, headers = 429, 'rate limit exceeded', {'Retry-After': '7'}
    class Flaky:
        n = 0
        def get(self, url, params=None, timeout=None):
            Flaky.n += 1
            return Resp429() if Flaky.n == 1 else _Resp({'ok': 1})
    clock['slept'].clear()
    assert cc._get_json(Flaky(), 'http://x/rest/api/content/1') == {'ok': 1}
    assert 7.0 in clock['slept']         # Retry-After만큼 쉬고 재시도


def test_status_view_renders_in_every_state(monkeypatch, tmp_path):
    """진행 중 표시가 컴포넌트 인자 오류(dbc.Spinner에 className 등)로 예외를 내면 Dash 콜백이
    조용히 실패해 화면에 아무것도 안 나온다 — 모든 상태가 예외 없이 그려지는지 확인."""
    import time
    from services import web_pipeline_runner as wpr
    monkeypatch.setattr(wpr, 'WEB_UPDATES_DIR', str(tmp_path))
    import app  # noqa: F401  (페이지 등록 후 컴포넌트 모듈 import)
    import components.admin_data_update_tab as tab
    for state in (dict(status='idle'),
                  dict(status='running', message='제목', count=3, root='123', started_at=time.time() - 5),
                  dict(status='done', count=2, root='123', finished_at='t'),
                  dict(status='error', message='boom')):
        confl_tree._save_state(**state)
        assert tab._confl_tree_status_view() is not None
