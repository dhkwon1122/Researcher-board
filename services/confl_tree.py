"""
컨플루언스 하위 페이지 목록 추출(2026-10, 사용자 요청) — "데이터 업데이트" 탭 하단.

상위 페이지(페이지 ID 또는 주소)를 입력하면 그 아래 모든 하위 페이지(자식 → 손자 →
최하위까지, 또는 지정한 단계까지)의 제목과 페이지 ID(10자리)를 엑셀로 내려받게 한다.
과제별컨플의 컨플 주소(페이지 ID)를 일일이 확인하던 작업을 대신한다.

하위 페이지가 많으면 요청 시간이 길어질 수 있어(gunicorn 워커 타임아웃) 백그라운드
스레드로 실행하고 화면은 진행 상황을 폴링한다. 동시에 하나의 작업만 허용한다(결과는
메모리에만 보관, 파일로 남기지 않음 — 서버에 페이지 목록 사본을 쌓지 않기 위해).
"""

import io
import os
import sys
import threading
from datetime import datetime

_PIPELINE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'pipeline')
sys.path.insert(0, os.path.abspath(_PIPELINE_DIR))

import confluence_client  # noqa: E402

_lock = threading.Lock()
_state: dict = {'status': 'idle', 'message': '', 'count': 0, 'rows': None, 'root': '', 'finished_at': ''}

EXCEL_HEADERS = ['단계', '제목', '페이지 ID', '상위 페이지 제목', '상위 페이지 ID', '경로']


def snapshot() -> dict:
    with _lock:
        return {k: v for k, v in _state.items() if k != 'rows'} | {'has_result': bool(_state['rows'])}


def is_running() -> bool:
    with _lock:
        return _state['status'] == 'running'


def start(root_address: str, max_depth: int | None) -> tuple[bool, str]:
    """작업 시작. 이미 실행 중이면 (False, 사유)."""
    root = confluence_client.normalize_address(root_address)
    if not root:
        return False, '상위 페이지 주소(또는 페이지 ID)를 입력하세요.'
    with _lock:
        if _state['status'] == 'running':
            return False, '이미 하위 페이지를 추출하는 중입니다. 끝난 뒤 다시 시도하세요.'
        _state.update(status='running', message='시작', count=0, rows=None, root=root, finished_at='')
    threading.Thread(target=_run, args=(root, max_depth), daemon=True).start()
    return True, ''


def _run(root: str, max_depth: int | None) -> None:
    def _progress(count: int, title: str) -> None:
        with _lock:
            _state['count'] = count
            _state['message'] = title

    try:
        rows = confluence_client.crawl_descendants(root, max_depth=max_depth, progress=_progress)
        with _lock:
            _state.update(status='done', rows=rows, count=len(rows), message='완료',
                          finished_at=datetime.now().strftime('%Y-%m-%d %H:%M:%S'))
    except Exception as exc:  # noqa: BLE001 — 사유를 화면에 그대로 보여줘야 함
        with _lock:
            _state.update(status='error', rows=None, message=str(exc)[:500],
                          finished_at=datetime.now().strftime('%Y-%m-%d %H:%M:%S'))


def build_workbook_bytes(rows: list[dict]) -> bytes:
    """결과 행을 엑셀(xlsx) 바이트로. 페이지 ID는 '3862782334'처럼 숫자 서식이 아닌
    텍스트 셀로 써서 엑셀이 지수/실수로 바꾸지 않게 한다."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = '하위 페이지'
    ws.append(EXCEL_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for r in rows:
        ws.append([r['depth'], r['title'], str(r['id']), r['parent_title'], str(r['parent_id']), r['path']])
    for row in ws.iter_rows(min_row=2):
        for idx in (2, 4):   # 페이지 ID 열들을 텍스트 서식으로
            row[idx].number_format = '@'
    widths = [6, 50, 14, 40, 14, 80]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'A2'
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def result_workbook() -> tuple[str, bytes] | None:
    """(파일명, 바이트) — 결과가 없으면 None."""
    with _lock:
        rows = _state['rows']
        root = _state['root']
    if not rows:
        return None
    return f'컨플_하위페이지_{root}_{datetime.now().strftime("%Y%m%d_%H%M")}.xlsx', build_workbook_bytes(rows)
