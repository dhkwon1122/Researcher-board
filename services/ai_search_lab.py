"""
AI 검색 테스트 랩(2026-10, 사용자 요청) — 관리자 "AI 검색 테스트" 탭의 백엔드.

보유 데이터를 바탕으로 내부 LLM이 사전 질문을 만들고, 그 질문을 실제 AI 검색 파이프라인
(services/nl_query: parse_question → execute_query → 결과 설명)에 그대로 넣어 답을 확인하고,
품질을 세 겹으로 평가해 질문/규칙을 개선하는 데 쓴다.

  1) 질문 생성   : 카테고리별로 LLM이 질문을 만든다(재료 = 부서·강점 분야·전공·직급 같은 용어 목록).
                   + 정답을 데이터에서 직접 계산할 수 있는 "정답 대조" 질문(골든)을 코드로 만든다.
  2) 일괄 실행   : 질문을 하나씩 실행(실제 검색과 같은 경로, 일반 AI 검색 로그에는 남기지 않음).
  3) 평가        : (a) 기계적 검사(오류·미지원·0건·느림·잘림)
                   (b) 정답 대조(골든 질문은 사번 집합의 정밀도/재현율/F1)
                   (c) LLM 심사(1~5점 + 문제 유형 + 이유)
  4) 개선        : 실행끼리 점수 비교, 저점수 묶음에 대한 "규칙 설정" 제안문 생성(자동 적용은 하지 않음).

오래 걸릴 수 있어 백그라운드 스레드로 돌리고 상태/결과는 파일(data/web_updates/ai_search_lab/)에
둔다 — gunicorn 워커가 여럿이라 프로세스 메모리에만 두면 화면 폴링이 다른 워커로 갈 때 진행이
안 보인다(services/confl_tree.py와 같은 이유). 일괄 실행은 시작한 관리자의 권한으로 수행한다
(services.auth.acting_as) — 권한이 필요한 테이블은 그 관리자가 볼 수 있는 범위에서만 조회된다.
"""

import io
import json
import os
import random
import re
import sys
import threading
import time
from datetime import datetime

_PIPELINE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'pipeline')
sys.path.insert(0, os.path.abspath(_PIPELINE_DIR))

import llm_client  # noqa: E402

from services import auth, data_store  # noqa: E402

CATEGORIES = {
    '전문성 검색': '특정 기술·연구 분야(강점 분야/키워드)를 가진 연구원 찾기. 예: "강화학습 전문가를 찾아줘"',
    '학력·나이·직급 조건': '학위, 전공, 나이대, 직급(CL) 같은 인사 기본 조건. 예: "물리학 박사 연구원"',
    '유사 연구원': '특정 연구원과 비슷한 전문성을 가진 사람 찾기(실제 존재하는 사번/이름이 필요하면 만들지 말고 "OO 연구원" 같은 자리표시자를 쓰지 말 것 — 이 카테고리는 이름 대신 분야 설명으로 질문)',
    '논문·특허': '논문 편수, 저널, 특허 건수·등급, 대표발명자 같은 실적 조건',
    '조직·과제': '부서·과제·팀 소속, 과제 참여 이력, 조직별 인원',
    '복합 조건': '위 조건 2~3개를 조합한 질문. 예: "AI 분야이면서 박사이고 특허가 2건 이상인 연구원"',
    '동의어·오타·모호한 표현': '구어체, 줄임말, 오타, 영문/한글 혼용, 애매한 표현으로 같은 의도를 묻는 질문',
    '지원하지 않는 질문': '이 시스템 데이터로는 답할 수 없는 질문(예: 날씨, 개인 연락처, 연봉 상세, 데이터에 없는 항목). 정중히 거절해야 정답',
}
GOLDEN_CATEGORY = '정답 대조'
NEGATIVE_CATEGORY = '지원하지 않는 질문'

ISSUE_TYPES = ['정확', '조건 누락', '조건 오해석', '컬럼/표 선택 오류', '결과 과다', '결과 없음', '거절해야 하는데 답함',
               '거절이 부적절', '기타']

_STALE_SECONDS = 900          # 'running' 작업이 이만큼 갱신이 없으면 중단된 것으로 본다(LLM 한 번이 길 수 있음)
_SLOW_SECONDS = 60
_SAMPLE_ROWS = 10
_JUDGE_ROWS = 10
_lock = threading.Lock()


# ── 저장소(파일) ─────────────────────────────────────────────────────────────

def _dir(*parts) -> str:
    from services import web_pipeline_runner as wpr    # WEB_UPDATES_DIR을 호출 시점에 읽는다(테스트에서 교체 가능)
    d = os.path.join(wpr.WEB_UPDATES_DIR, 'ai_search_lab', *parts)
    os.makedirs(d, exist_ok=True)
    return d


def _write_json(path: str, data) -> None:
    tmp = f'{path}.{os.getpid()}.{threading.get_ident()}.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(data, f, ensure_ascii=False)
    try:
        os.chmod(tmp, 0o600)
    except OSError:
        pass
    os.replace(tmp, path)


def _read_json(path: str, default):
    try:
        with open(path, encoding='utf-8') as f:
            return json.load(f)
    except (OSError, ValueError):
        return default


def _safe_name(name: str) -> str:
    return re.sub(r'[^0-9A-Za-z가-힣_\-. ]+', '_', (name or '').strip())[:60].strip() or 'set'


# ── 작업(job) 상태 ───────────────────────────────────────────────────────────

def _job_path() -> str:
    return os.path.join(_dir(), 'job.json')


def _stop_path() -> str:
    return os.path.join(_dir(), 'stop.flag')


def job_state() -> dict:
    base = {'kind': '', 'status': 'idle', 'message': '', 'done': 0, 'total': 0, 'run_id': '',
            'started_at': 0.0, 'updated_at': 0.0}
    base.update(_read_json(_job_path(), {}))
    if base['status'] == 'running' and time.time() - float(base['updated_at'] or 0) > _STALE_SECONDS:
        base.update(status='error', message='작업이 응답 없이 중단됐습니다(서버 재시작 등). 다시 실행하세요.')
    base['elapsed'] = int(time.time() - float(base['started_at'])) if base['status'] == 'running' and base['started_at'] else 0
    return base


def is_busy() -> bool:
    return job_state()['status'] == 'running'


def _save_job(**updates) -> dict:
    st = job_state()
    st.pop('elapsed', None)
    st.update(updates)
    st['updated_at'] = time.time()
    _write_json(_job_path(), st)
    return st


def _begin(kind: str, total: int, message: str, run_id: str = '') -> tuple[bool, str]:
    with _lock:
        if is_busy():
            return False, '다른 테스트 작업이 실행 중입니다. 끝나거나 중지한 뒤 다시 시도하세요.'
        try:
            os.remove(_stop_path())
        except OSError:
            pass
        _save_job(kind=kind, status='running', message=message, done=0, total=total, run_id=run_id,
                  started_at=time.time())
    return True, ''


def stop() -> None:
    """실행 중인 일괄 실행에 중지 요청(현재 질문까지 마치고 멈춘다)."""
    with open(_stop_path(), 'w') as f:
        f.write('1')


def _stop_requested() -> bool:
    return os.path.exists(_stop_path())


# ── 질문 목록(초안/세트) ──────────────────────────────────────────────────────

def load_draft() -> list[dict]:
    return _read_json(os.path.join(_dir(), 'draft.json'), [])


def save_draft(items: list[dict]) -> None:
    _write_json(os.path.join(_dir(), 'draft.json'), _normalize_items(items))


def _normalize_items(items: list[dict]) -> list[dict]:
    out, seen = [], set()
    for i, it in enumerate(items or []):
        q = str((it or {}).get('question') or '').strip()
        if not q or q in seen:
            continue
        seen.add(q)
        out.append({'id': str(it.get('id') or f'q{i + 1}'), 'category': str(it.get('category') or '기타'),
                    'question': q, 'expected': it.get('expected') or None})
    return out


def list_sets() -> list[str]:
    return sorted(os.path.splitext(f)[0] for f in os.listdir(_dir('sets')) if f.endswith('.json'))


def save_set(name: str, items: list[dict]) -> str:
    name = _safe_name(name)
    _write_json(os.path.join(_dir('sets'), f'{name}.json'), _normalize_items(items))
    return name


def load_set(name: str) -> list[dict]:
    return _read_json(os.path.join(_dir('sets'), f'{_safe_name(name)}.json'), [])


# ── 질문 생성 ────────────────────────────────────────────────────────────────

def _top_values(values, n: int) -> list[str]:
    counts: dict = {}
    for v in values:
        v = str(v or '').strip()
        if v and v.lower() != 'nan':
            counts[v] = counts.get(v, 0) + 1
    return [v for v, _ in sorted(counts.items(), key=lambda kv: -kv[1])[:n]]


def build_vocab() -> dict:
    """질문 생성 재료 — 조직/분야 용어만 모은다(개인 이름·평가 정보는 넣지 않는다)."""
    researchers = data_store.read_processed('researchers')
    education = data_store.read_processed('education')
    vocab = {'부서': [], '직급': [], '강점 분야': [], '강점 키워드': [], '전공': [], '학위': [], '과제': []}
    if not researchers.empty:
        if 'department' in researchers.columns:
            vocab['부서'] = _top_values(researchers['department'], 30)
        if 'position' in researchers.columns:
            vocab['직급'] = _top_values(researchers['position'], 12)
    if not education.empty:
        if 'major' in education.columns:
            vocab['전공'] = _top_values(education['major'], 30)
        if 'degree' in education.columns:
            vocab['학위'] = _top_values(education['degree'], 6)
    fields, kws = [], []
    for p in data_store.read_expertise_profiles().values():
        fields += list(p.get('strength_fields') or [])
        kws += list(p.get('strength_keywords') or [])
    vocab['강점 분야'] = _top_values(fields, 40)
    vocab['강점 키워드'] = _top_values(kws, 40)
    tasks = data_store.read_processed('tasks')
    if not tasks.empty and 'task_name' in tasks.columns:
        vocab['과제'] = _top_values(tasks['task_name'], 25)
    return vocab


_GEN_SYSTEM = """# Role
당신은 사내 연구원 검색 시스템("AI 검색")의 품질 테스트용 질문을 만드는 QA 담당자입니다.

# Guidelines
1. 실제 사용자(HR 담당자, 부서장)가 자연어로 입력할 법한 짧은 한국어 질문을 만드세요.
2. 주어진 [데이터 용어]에 실제로 있는 값을 섞어 쓰되, 매번 같은 값만 쓰지 말고 다양하게 쓰세요.
3. 서로 표현·난이도·조건 개수가 다르게 만드세요. 같은 질문을 반복하지 마세요.
4. 반드시 아래 JSON으로만 답하세요: {"questions": ["질문1", "질문2", ...]}
"""


def _parse_questions(raw: str) -> list[str]:
    try:
        data = json.loads(llm_client.extract_json(raw))
    except (json.JSONDecodeError, TypeError):
        return []
    qs = data.get('questions') if isinstance(data, dict) else data
    return [str(q).strip() for q in (qs or []) if str(q).strip()]


def generate_llm_questions(categories: list[str], per_category: int, vocab: dict | None = None,
                           progress=None) -> list[dict]:
    vocab = vocab or build_vocab()
    vocab_text = '\n'.join(f'- {k}: {", ".join(v[:25])}' for k, v in vocab.items() if v) or '(데이터 용어 없음)'
    items = []
    for ci, cat in enumerate(categories):
        if _stop_requested():
            break
        prompt = (f'[카테고리] {cat}\n[설명] {CATEGORIES.get(cat, "")}\n[만들 개수] {per_category}\n\n'
                  f'[데이터 용어]\n{vocab_text}\n')
        raw = llm_client.call_llm(prompt, _GEN_SYSTEM, temperature=0.8, max_tokens=2500)
        for q in _parse_questions(raw)[:per_category]:
            items.append({'id': '', 'category': cat, 'question': q, 'expected': None})
        if progress:
            progress(ci + 1, len(categories), cat)
    return items


# ── 정답 대조(골든) 질문 ─────────────────────────────────────────────────────

def _current_ids(researchers) -> set:
    if researchers.empty:
        return set()
    df = researchers
    if 'is_current' in df.columns:
        df = df[df['is_current'].astype(str).str.upper() != 'N']
    return set(df['researcher_id'].astype(str).str.zfill(8))


def golden_items(max_each: int = 3, seed: int = 7) -> list[dict]:
    """데이터에서 정답 사번 집합을 직접 계산할 수 있는 질문. 기대 결과가 비어 있거나 너무 큰 질문은 뺀다."""
    rng = random.Random(seed)
    researchers = data_store.read_processed('researchers')
    current = _current_ids(researchers)
    if not current:
        return []
    items = []

    def add(question: str, ids: set, desc: str) -> None:
        ids = set(ids) & current
        if 1 <= len(ids) <= max(30, len(current) // 2):
            items.append({'id': '', 'category': GOLDEN_CATEGORY, 'question': question,
                          'expected': {'ids': sorted(ids), 'desc': f'{desc} — 기대 {len(ids)}명'}})

    education = data_store.read_processed('education')
    if not education.empty and 'degree' in education.columns:
        edu = education.assign(rid=education['researcher_id'].astype(str).str.zfill(8))
        add('박사 학위를 가진 연구원을 찾아줘', set(edu[edu['degree'].astype(str).str.contains('박사')]['rid']),
            '학력에 박사 포함')
        if 'major' in edu.columns:
            for major in _top_values(edu['major'], 10)[:max_each]:
                add(f'전공이 {major}인 연구원', set(edu[edu['major'].astype(str) == major]['rid']), f'전공={major}')

    if not researchers.empty:
        r = researchers.assign(rid=researchers['researcher_id'].astype(str).str.zfill(8))
        if 'department' in r.columns:
            for dept in _top_values(r[r['rid'].isin(current)]['department'], 8)[:max_each]:
                add(f'{dept} 소속 연구원 명단', set(r[r['department'].astype(str) == dept]['rid']), f'부서={dept}')
        if 'position' in r.columns:
            for pos in _top_values(r[r['rid'].isin(current)]['position'], 6)[:2]:
                add(f'직급이 {pos}인 연구원', set(r[r['position'].astype(str) == pos]['rid']), f'직급={pos}')

    pubs = data_store.read_processed('publications')
    if not pubs.empty:
        cnt = pubs.assign(rid=pubs['researcher_id'].astype(str).str.zfill(8)).groupby('rid').size()
        for k in sorted({max(2, int(cnt.median())), max(3, int(cnt.quantile(0.75)))}):
            add(f'논문을 {k}편 이상 쓴 연구원', set(cnt[cnt >= k].index), f'논문 {k}편 이상')

    pats = data_store.read_processed('patents')
    if not pats.empty and 'application_id' in pats.columns:
        cnt = (pats.assign(rid=pats['researcher_id'].astype(str).str.zfill(8))
               .drop_duplicates(['rid', 'application_id']).groupby('rid').size())
        for k in sorted({max(2, int(cnt.median())), max(3, int(cnt.quantile(0.75)))}):
            add(f'특허가 {k}건 이상인 연구원', set(cnt[cnt >= k].index), f'특허 {k}건 이상')

    profiles = data_store.read_expertise_profiles()
    if profiles:
        fields = _top_values([f for p in profiles.values() for f in (p.get('strength_fields') or [])], 10)
        for field in fields[:max_each]:
            ids = {str(rid).zfill(8) for rid, p in profiles.items() if field in (p.get('strength_fields') or [])}
            add(f'{field} 분야 전문가를 찾아줘', ids, f'강점 분야에 "{field}" 포함(의미 유사 표기는 기대에 없음)')
    rng.shuffle(items)
    return items


# ── 실행 + 평가 ──────────────────────────────────────────────────────────────

_JUDGE_SYSTEM = """# Role
당신은 사내 연구원 검색("AI 검색") 결과의 품질을 심사하는 QA 심사위원입니다.

# Guidelines
1. 질문의 의도와 조건을 모두 따져, 시스템이 만든 조회(SQL 또는 검색 유형)와 결과 표가 그에 맞는지 판단하세요.
2. score: 5=조건과 결과가 정확히 맞음 / 4=사소한 누락·과다 / 3=일부 조건 누락이나 결과가 애매함 /
   2=상당 부분 틀림 / 1=전혀 다르거나 오류.
3. 카테고리가 "지원하지 않는 질문"이면 정중히 거절하거나 결과를 내지 않는 것이 정답입니다(결과를 내놓으면 낮은 점수).
4. 결과 표에 없는 사실을 추측하지 말고 주어진 정보만으로 판단하세요.
5. issue_type은 다음 중 하나: 정확, 조건 누락, 조건 오해석, 컬럼/표 선택 오류, 결과 과다, 결과 없음,
   거절해야 하는데 답함, 거절이 부적절, 기타
6. reason은 한국어 1~2문장. 반드시 JSON으로만 답하세요: {"score": 1~5, "issue_type": "...", "reason": "..."}
"""


def _result_ids(result: dict) -> set:
    cols = [str(c) for c in (result.get('columns') or [])]
    if 'researcher_id' not in cols:
        return set()
    idx = cols.index('researcher_id')
    return {str(r[idx]).strip().zfill(8) for r in (result.get('rows') or []) if r and r[idx] not in (None, '')}


def _golden_compare(expected_ids: list, got: set) -> dict:
    exp = set(expected_ids)
    tp = len(exp & got)
    precision = tp / len(got) if got else 0.0
    recall = tp / len(exp) if exp else 0.0
    f1 = (2 * precision * recall / (precision + recall)) if (precision + recall) else 0.0
    return {'expected': len(exp), 'returned': len(got), 'precision': round(precision, 3),
            'recall': round(recall, 3), 'f1': round(f1, 3),
            'missing_sample': sorted(exp - got)[:5], 'extra_sample': sorted(got - exp)[:5]}


def _judge(item: dict, result: dict, answer: str) -> dict | None:
    cols = result.get('labels') or result.get('columns') or []
    rows = (result.get('rows') or [])[:_JUDGE_ROWS]
    table = ' | '.join(str(c) for c in cols) + '\n' + '\n'.join(
        ', '.join('' if v is None else str(v) for v in r) for r in rows)
    prompt = (f'[질문] {item["question"]}\n[카테고리] {item.get("category", "")}\n'
              f'[분류된 유형] {result.get("intent", "")}\n[생성된 SQL]\n{result.get("sql", "(없음)")}\n'
              f'[결과: 총 {result.get("total_rows", len(result.get("rows") or []))}건 중 상위 {len(rows)}건]\n{table}\n'
              f'[시스템 안내/설명]\n{result.get("note", "")}\n{answer}\n')
    raw = llm_client.call_llm(prompt, _JUDGE_SYSTEM, temperature=0.0, max_tokens=500)
    if not raw:
        return None
    try:
        data = json.loads(llm_client.extract_json(raw))
        score = int(data.get('score'))
    except (json.JSONDecodeError, TypeError, ValueError):
        return None
    issue = str(data.get('issue_type') or '기타')
    return {'score': max(1, min(5, score)), 'issue_type': issue if issue in ISSUE_TYPES else '기타',
            'reason': str(data.get('reason') or '')[:300]}


def _status(item: dict, rec: dict) -> str:
    negative = item.get('category') == NEGATIVE_CATEGORY
    flags = rec['flags']
    judge = rec.get('judge')
    golden = rec.get('golden')
    if negative:
        if rec['intent'] in ('unsupported', 'error') or rec['row_count'] == 0:
            return '양호' if not (judge and judge['score'] <= 2) else '주의'
        return '실패' if rec['row_count'] > 0 and not (judge and judge['score'] >= 4) else '주의'
    if 'error' in flags or 'unsupported' in flags:
        return '실패'
    if golden and golden['f1'] < 0.5:
        return '실패'
    if judge and judge['score'] <= 2:
        return '실패'
    if 'empty' in flags or 'slow' in flags or 'truncated' in flags:
        return '주의'
    if golden and golden['f1'] < 0.8:
        return '주의'
    if judge and judge['score'] == 3:
        return '주의'
    return '양호'


def run_one(item: dict, current_only: bool = True, period: tuple | None = None, judge: bool = True) -> dict:
    """질문 하나를 실제 AI 검색 경로로 실행하고 평가 결과 record를 만든다(일반 검색 로그에는 남기지 않음)."""
    from services import nl_query
    t0 = time.time()
    rec = {'id': item.get('id', ''), 'category': item.get('category', ''), 'question': item['question'],
           'intent': '', 'sql': '', 'columns': [], 'sample_rows': [], 'row_count': 0, 'total_rows': 0,
           'answer': '', 'note': '', 'seconds': 0.0, 'flags': [], 'golden': None, 'judge': None, 'status': ''}
    result = {}
    try:
        parsed = nl_query.parse_question(item['question'])
        result = nl_query.execute_query(parsed, current_only=current_only, period=period)
        if result.get('intent') not in ('error', 'unsupported'):
            rec['answer'] = nl_query._generate_answer_summary(item['question'], result)
    except Exception as exc:  # noqa: BLE001 — 한 질문의 실패가 일괄 실행을 멈추지 않게
        result = {'intent': 'error', 'note': f'{type(exc).__name__}: {exc}', 'columns': [], 'rows': []}
    rec['seconds'] = round(time.time() - t0, 1)
    rec.update(intent=str(result.get('intent', '')), sql=str(result.get('sql', '') or ''),
               columns=list(result.get('labels') or result.get('columns') or []),
               sample_rows=[[None if v is None else str(v) for v in r] for r in (result.get('rows') or [])[:_SAMPLE_ROWS]],
               row_count=len(result.get('rows') or []),
               total_rows=int(result.get('total_rows') or len(result.get('rows') or [])),
               note=str(result.get('note', '') or ''))

    flags = rec['flags']
    if rec['intent'] == 'error':
        flags.append('error')
    elif rec['intent'] == 'unsupported':
        flags.append('unsupported')
    elif rec['row_count'] == 0:
        flags.append('empty')
    if rec['seconds'] > _SLOW_SECONDS:
        flags.append('slow')
    if rec['total_rows'] > rec['row_count'] > 0:
        flags.append('truncated')
    if rec['row_count'] and not rec['answer']:
        flags.append('no_answer_text')

    exp = item.get('expected')
    if exp and exp.get('ids') is not None:
        rec['golden'] = _golden_compare(exp['ids'], _result_ids(result))
        rec['golden']['desc'] = exp.get('desc', '')
    if judge and rec['intent'] != 'error':
        try:
            rec['judge'] = _judge(item, result, rec['answer'])
        except Exception:  # noqa: BLE001 — 심사 실패는 점수만 비운다
            rec['judge'] = None
    rec['status'] = _status(item, rec)
    return rec


def _run_path(run_id: str) -> str:
    return os.path.join(_dir('runs'), f'{run_id}.json')


def start_run(items: list[dict], user: dict | None, label: str = '', current_only: bool = True,
              period: tuple | None = None, judge: bool = True) -> tuple[bool, str, str]:
    """일괄 실행 시작 → (성공여부, 사유, run_id)."""
    items = _normalize_items(items)
    if not items:
        return False, '실행할 질문이 없습니다.', ''
    run_id = datetime.now().strftime('%Y%m%d_%H%M%S')
    ok, reason = _begin('run', len(items), '시작', run_id)
    if not ok:
        return False, reason, ''
    run = {'run_id': run_id, 'label': label or run_id, 'started': datetime.now().strftime('%Y-%m-%d %H:%M:%S'),
           'finished': '', 'options': {'current_only': current_only, 'period': list(period) if period else None,
                                       'judge': judge}, 'status': 'running', 'results': [], 'suggestion': ''}
    _write_json(_run_path(run_id), run)
    threading.Thread(target=_run_loop, args=(run, items, user, current_only, period, judge), daemon=True).start()
    return True, '', run_id


def _run_loop(run: dict, items: list[dict], user, current_only, period, judge) -> None:
    try:
        with auth.acting_as(user or {'user_id': 'ai-search-lab', 'role': 'admin', 'is_admin': True,
                                       'permissions': {}, 'eval_excluded_dep_ids': []}):
            for i, item in enumerate(items):
                if _stop_requested():
                    run['status'] = 'stopped'
                    break
                _save_job(message=item['question'][:60], done=i)
                run['results'].append(run_one(item, current_only, period, judge))
                _write_json(_run_path(run['run_id']), run)
        if run['status'] == 'running':
            run['status'] = 'done'
        run['finished'] = datetime.now().strftime('%Y-%m-%d %H:%M:%S')
        run['summary'] = summarize(run)
        _write_json(_run_path(run['run_id']), run)
        _save_job(status='done' if run['status'] == 'done' else 'stopped', done=len(run['results']), message='완료')
    except Exception as exc:  # noqa: BLE001
        run['status'] = 'error'
        _write_json(_run_path(run['run_id']), run)
        _save_job(status='error', message=f'{type(exc).__name__}: {exc}'[:300])


def summarize(run: dict) -> dict:
    res = run.get('results') or []
    judged = [r['judge']['score'] for r in res if r.get('judge')]
    f1s = [r['golden']['f1'] for r in res if r.get('golden')]
    by_cat: dict = {}
    for r in res:
        c = by_cat.setdefault(r['category'], {'n': 0, '양호': 0, '주의': 0, '실패': 0})
        c['n'] += 1
        c[r['status']] = c.get(r['status'], 0) + 1
    return {
        'n': len(res), '양호': sum(r['status'] == '양호' for r in res), '주의': sum(r['status'] == '주의' for r in res),
        '실패': sum(r['status'] == '실패' for r in res),
        'avg_score': round(sum(judged) / len(judged), 2) if judged else None,
        'avg_f1': round(sum(f1s) / len(f1s), 3) if f1s else None,
        'avg_seconds': round(sum(r['seconds'] for r in res) / len(res), 1) if res else None,
        'by_category': by_cat,
    }


def list_runs() -> list[dict]:
    out = []
    for f in sorted(os.listdir(_dir('runs')), reverse=True):
        if not f.endswith('.json'):
            continue
        run = _read_json(os.path.join(_dir('runs'), f), {})
        if run:
            out.append({'run_id': run.get('run_id', f[:-5]), 'label': run.get('label', ''),
                        'started': run.get('started', ''), 'status': run.get('status', ''),
                        'summary': run.get('summary') or summarize(run)})
    return out


def load_run(run_id: str) -> dict:
    return _read_json(_run_path(run_id), {})


def compare_runs(run_a: dict, run_b: dict) -> list[dict]:
    """질문 기준으로 두 실행(a=이전, b=이후)의 점수/상태 변화."""
    a = {r['question']: r for r in run_a.get('results') or []}
    rows = []
    for r in run_b.get('results') or []:
        prev = a.get(r['question'])
        score_b = r['judge']['score'] if r.get('judge') else None
        score_a = prev['judge']['score'] if prev and prev.get('judge') else None
        f1_b = r['golden']['f1'] if r.get('golden') else None
        f1_a = prev['golden']['f1'] if prev and prev.get('golden') else None
        rows.append({'question': r['question'], 'category': r['category'],
                     'status_before': prev['status'] if prev else '', 'status_after': r['status'],
                     'score_before': score_a, 'score_after': score_b,
                     'f1_before': f1_a, 'f1_after': f1_b,
                     'delta_score': (score_b - score_a) if (score_a is not None and score_b is not None) else None})
    return rows


# ── 개선 제안 + 생성 작업 ─────────────────────────────────────────────────────

_SUGGEST_SYSTEM = """# Role
당신은 사내 연구원 검색("AI 검색")의 프롬프트 개선 담당자입니다. QA에서 문제로 나온 질문들을 보고,
관리자가 "규칙 설정"에 추가할 수 있는 짧은 규칙 문장을 제안합니다.

# Guidelines
1. 문제 유형이 반복되는 것부터 묶어서, 일반화할 수 있는 규칙 3~8개를 제안하세요(특정 질문 하나에만 맞는 규칙은 피하세요).
2. 각 규칙은 "- "로 시작하는 한 줄 한국어 문장(용어 정의, 동의어, 조건 해석 방법, 출력 형식 지시)으로 쓰세요.
3. 기존 규칙과 중복되거나 모순되는 내용은 쓰지 마세요.
4. 마지막에 "# 근거" 소제목 아래 어떤 질문 유형 때문에 제안했는지 2~4줄로 적으세요.
"""


def start_suggest(run_id: str) -> tuple[bool, str]:
    run = load_run(run_id)
    problems = [r for r in run.get('results') or [] if r.get('status') != '양호'][:25]
    if not problems:
        return False, '개선이 필요한(양호가 아닌) 질문이 없습니다.'
    ok, reason = _begin('suggest', 1, '개선 제안 생성 중', run_id)
    if not ok:
        return False, reason

    def _work():
        from services import query_settings
        try:
            lines = []
            for r in problems:
                j = r.get('judge') or {}
                lines.append(f'- [{r["category"]}] {r["question"]} → 유형 {r["intent"]}, {r["row_count"]}건, '
                             f'상태 {r["status"]}, 문제 {j.get("issue_type", "")}: {j.get("reason", "")}'
                             + (f' / SQL: {r["sql"][:200]}' if r.get('sql') else ''))
            prompt = f'[현재 규칙 설정]\n{query_settings.read_rules() or "(없음)"}\n\n[문제 질문들]\n' + '\n'.join(lines)
            raw = llm_client.call_llm(prompt, _SUGGEST_SYSTEM, temperature=0.3, max_tokens=1500)
            run['suggestion'] = raw.strip() if raw else '(LLM 응답이 없어 제안을 만들지 못했습니다.)'
            _write_json(_run_path(run_id), run)
            _save_job(status='done', done=1, message='개선 제안 완료')
        except Exception as exc:  # noqa: BLE001
            _save_job(status='error', message=f'{type(exc).__name__}: {exc}'[:300])

    threading.Thread(target=_work, daemon=True).start()
    return True, ''


def start_generate(categories: list[str], per_category: int, include_golden: bool, replace: bool,
                   user: dict | None) -> tuple[bool, str]:
    categories = [c for c in categories if c in CATEGORIES]
    if not categories and not include_golden:
        return False, '생성할 카테고리를 하나 이상 고르세요.'
    ok, reason = _begin('generate', max(1, len(categories)), '질문 생성 중')
    if not ok:
        return False, reason

    def _work():
        try:
            with auth.acting_as(user or {'user_id': 'ai-search-lab', 'role': 'admin', 'is_admin': True,
                                           'permissions': {}, 'eval_excluded_dep_ids': []}):
                new_items = []
                if include_golden:
                    new_items += golden_items()
                new_items += generate_llm_questions(
                    categories, per_category,
                    progress=lambda d, t, c: _save_job(done=d, message=f'{c} 생성 완료'))
                base = [] if replace else load_draft()
                save_draft(base + new_items)
            _save_job(status='done', done=len(categories), message=f'질문 {len(new_items)}개 추가')
        except Exception as exc:  # noqa: BLE001
            _save_job(status='error', message=f'{type(exc).__name__}: {exc}'[:300])

    threading.Thread(target=_work, daemon=True).start()
    return True, ''


# ── 엑셀 ─────────────────────────────────────────────────────────────────────

def build_run_workbook(run: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = '결과'
    head = ['No', '카테고리', '질문', '상태', '유형(intent)', '건수', '전체건수', '시간(초)', '정답 대조 F1', '정밀도',
            '재현율', '심사 점수', '문제 유형', '심사 이유', '플래그', 'SQL', '시스템 설명/안내', '결과 일부(상위 5행)']
    ws.append(head)
    for r in run.get('results') or []:
        g, j = r.get('golden') or {}, r.get('judge') or {}
        sample = '\n'.join(', '.join(row) for row in (r.get('sample_rows') or [])[:5])
        ws.append([len(ws['A']), r['category'], r['question'], r['status'], r['intent'], r['row_count'],
                   r['total_rows'], r['seconds'], g.get('f1'), g.get('precision'), g.get('recall'),
                   j.get('score'), j.get('issue_type'), j.get('reason'), ', '.join(r['flags']), r.get('sql'),
                   (r.get('answer') or r.get('note') or ''), sample])
    for c in ws[1]:
        c.font = Font(bold=True)
    for i, w in enumerate([5, 16, 40, 8, 16, 7, 9, 8, 10, 8, 8, 8, 16, 40, 18, 50, 50, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.freeze_panes = 'D2'

    s = run.get('summary') or summarize(run)
    ws2 = wb.create_sheet('요약')
    ws2.append(['항목', '값'])
    for k, v in (('실행 이름', run.get('label')), ('시작', run.get('started')), ('종료', run.get('finished')),
                 ('질문 수', s['n']), ('양호', s['양호']), ('주의', s['주의']), ('실패', s['실패']),
                 ('평균 심사 점수', s['avg_score']), ('평균 정답 대조 F1', s['avg_f1']),
                 ('평균 응답 시간(초)', s['avg_seconds'])):
        ws2.append([k, v])
    ws2.append([])
    ws2.append(['카테고리', '질문 수', '양호', '주의', '실패'])
    for cat, c in s['by_category'].items():
        ws2.append([cat, c['n'], c.get('양호', 0), c.get('주의', 0), c.get('실패', 0)])
    if run.get('suggestion'):
        ws3 = wb.create_sheet('개선 제안')
        for line in run['suggestion'].splitlines():
            ws3.append([line])
        ws3.column_dimensions['A'].width = 120
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
