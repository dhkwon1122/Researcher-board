"""
추출한 컨플 하위 페이지(제목·페이지 ID)를 과제별컨플 엑셀의 "컨플 주소"에 자동 반영(2026-10, 사용자 요청).

흐름
  1) 추출 결과(services/confl_tree의 rows) 중 단계(depth)==PROJECT_DEPTH(2, 과제 페이지)만 후보로 쓴다.
     (단계 0 = 월별 페이지, 1 = 부서, 2 = 과제. 후보의 '상위 페이지 제목'이 부서명.)
  2) 서버에 마지막으로 올라간 과제별컨플 원본(data/web_updates/project_confl_address/)의 각 행에서
     "과제명"과 후보 제목을 정규화해 문자열 유사도로 비교한다. 정규화: '|' 뒤(월 표기) 제거,
     맨 앞 '[연구]' 같은 대괄호 꼬리표 제거, 공백·기호 제거, 소문자.
  3) 과제명 유사도 >= MIN_SIMILARITY(0.7)인 후보만 인정하고, 순위는 "과제명 유사도 80% + 소속↔부서명
     유사도 20%"로 매긴다(같은 이름의 과제가 여러 부서에 있을 때 구분). 한 페이지는 한 행에만 쓴다
     (점수 높은 쌍부터 배정).
  4) 매칭된 행은 "컨플 주소"를 페이지 ID(텍스트)로 덮어쓰고(이미 값이 있어도), 매칭 실패 행은 기존
     값을 건드리지 않는다. 마지막 열 뒤(최소 L열)에 "추출 결과"(성공/실패), "매칭된 제목", "유사도"
     열을 추가한다. 매칭되지 않은 후보 페이지는 "미매칭 페이지" 시트에 모은다.
  사용자는 이 엑셀을 확인한 뒤 데이터 업데이트의 과제별컨플로 다시 올린다.
"""

import difflib
import glob
import io
import os
import re
import sys

_PIPELINE_DIR = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'pipeline')
sys.path.insert(0, os.path.abspath(_PIPELINE_DIR))

PROJECT_DEPTH = 2
MIN_SIMILARITY = 0.7
TITLE_WEIGHT, DEPT_WEIGHT = 0.8, 0.2

HEADER_DEP = '소속'
HEADER_PROJECT = '과제명'
HEADER_ADDR = '컨플 주소'
HEADER_RESULT = '추출 결과'
HEADER_TITLE = '매칭된 제목'
HEADER_SCORE = '유사도'
MIN_RESULT_COL = 12   # L열


def normalize_title(text: str) -> str:
    """'[연구]2D Material' / "2D Material|'26.9월" 같은 표기 차이를 지운 비교용 문자열."""
    s = str(text or '').split('|', 1)[0].strip()
    s = re.sub(r'^(?:\[[^\]]*\]|\([^)]*\))\s*', '', s)      # 맨 앞 [연구]/(과제) 꼬리표 한 번
    s = re.sub(r'[\s\-_/·.,:;()\[\]{}\'"`~!@#$%^&*+=<>?]+', '', s)
    return s.lower()


def similarity(a: str, b: str) -> float:
    na, nb = normalize_title(a), normalize_title(b)
    if not na or not nb:
        return 0.0
    return difflib.SequenceMatcher(None, na, nb).ratio()


def _dept_similarity(a: str, b: str) -> float:
    na, nb = normalize_title(a), normalize_title(b)
    if not na or not nb:
        return 0.5          # 정보가 없으면 중립(순위에 영향 없음)
    return difflib.SequenceMatcher(None, na, nb).ratio()


def match(project_rows: list[dict], pages: list[dict]) -> dict[int, dict]:
    """project_rows: [{'idx', 'dept', 'project'}], pages: 추출 결과 중 과제 단계 행들.
    반환: {idx: {'page': page, 'title_sim', 'score'}} — 매칭된 행만. 한 페이지는 한 행에만 배정."""
    pairs = []
    for pr in project_rows:
        for pg in pages:
            title_sim = similarity(pr['project'], pg['title'])
            if title_sim < MIN_SIMILARITY:
                continue
            dept_sim = _dept_similarity(pr['dept'], pg.get('parent_title', ''))
            pairs.append((TITLE_WEIGHT * title_sim + DEPT_WEIGHT * dept_sim, title_sim, pr['idx'], pg))
    pairs.sort(key=lambda t: -t[0])
    used_rows, used_pages, result = set(), set(), {}
    for score, title_sim, idx, pg in pairs:
        if idx in used_rows or pg['id'] in used_pages:
            continue
        used_rows.add(idx)
        used_pages.add(pg['id'])
        result[idx] = {'page': pg, 'title_sim': title_sim, 'score': score}
    return result


def find_source_workbook() -> str | None:
    """서버에 마지막으로 올라간 과제별컨플 원본(xlsx/xlsb) 경로. 없으면 None."""
    from services import web_pipeline_runner as wpr
    files = [p for p in wpr.uploaded_files('project_confl_address')
             if p.lower().endswith(('.xlsx', '.xlsb'))]
    return max(files, key=os.path.getmtime) if files else None


def _load_workbook_from_source():
    """(openpyxl Workbook, 출처 설명). 업로드 원본(xlsx면 서식 보존, xlsb면 값만 xlsx로 변환)을 쓰고,
    없으면 처리된 project_confl_address.csv로 3열짜리 표를 만든다. 둘 다 없으면 ValueError."""
    from openpyxl import Workbook, load_workbook

    path = find_source_workbook()
    if path and path.lower().endswith('.xlsx'):
        return load_workbook(path), f'업로드 원본 {os.path.basename(path)}'
    if path:   # xlsb — pandas(pyxlsb)로 읽어 xlsx로 옮긴다(서식은 유지되지 않음)
        import pandas as pd
        df = pd.read_excel(path, header=0, engine='pyxlsb', dtype=str).fillna('')
        return _workbook_from_frame(df), f'업로드 원본 {os.path.basename(path)}(xlsb → xlsx 변환, 서식 제외)'

    from services.data_store import read_processed
    df = read_processed('project_confl_address')
    if df is not None and not df.empty:
        out = df.rename(columns={'dep_name': HEADER_DEP, 'project_name': HEADER_PROJECT,
                                 'confl_address': HEADER_ADDR})
        out = out[[c for c in (HEADER_DEP, HEADER_PROJECT, HEADER_ADDR) if c in out.columns]]
        return _workbook_from_frame(out.astype(str).replace('nan', '')), '처리된 project_confl_address 데이터(원본 파일 없음)'
    raise ValueError('업로드된 과제별컨플이 없습니다 — 먼저 데이터 업데이트에서 과제별컨플을 올려주세요.')


def _workbook_from_frame(df):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.append([str(c) for c in df.columns])
    for row in df.itertuples(index=False):
        ws.append(list(row))
    return wb


def apply_to_workbook(wb, rows: list[dict]) -> dict:
    """wb(과제별컨플)에 매칭 결과를 쓰고 요약을 반환한다. rows: 추출 결과 전체(단계 포함)."""
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    ws = wb.worksheets[0]
    headers = {str(c.value).strip(): c.column for c in ws[1] if c.value is not None and str(c.value).strip()}
    for needed in (HEADER_PROJECT, HEADER_ADDR):
        if needed not in headers:
            raise ValueError(f'과제별컨플 첫 행에 "{needed}" 헤더가 없습니다(현재 헤더: {list(headers)}).')
    col_dep = headers.get(HEADER_DEP)
    col_project, col_addr = headers[HEADER_PROJECT], headers[HEADER_ADDR]

    # 결과 열 위치 — 이미 "추출 결과"가 있으면(이전 반영본 재사용) 그 자리, 없으면 마지막 열 다음(최소 L열)
    if HEADER_RESULT in headers:
        col_result = headers[HEADER_RESULT]
    else:
        col_result = max(MIN_RESULT_COL, ws.max_column + 1)
    col_title, col_score = col_result + 1, col_result + 2
    bold = Font(bold=True)
    for col, text in ((col_result, HEADER_RESULT), (col_title, HEADER_TITLE), (col_score, HEADER_SCORE)):
        cell = ws.cell(row=1, column=col, value=text)
        cell.font = bold

    pages = [r for r in rows if r.get('depth') == PROJECT_DEPTH]
    project_rows = []
    for r in range(2, ws.max_row + 1):
        name = ws.cell(row=r, column=col_project).value
        if name is None or not str(name).strip():
            continue
        dept = ws.cell(row=r, column=col_dep).value if col_dep else ''
        project_rows.append({'idx': r, 'dept': str(dept or ''), 'project': str(name)})

    matched = match(project_rows, pages)
    fail_fill = PatternFill('solid', fgColor='FFF3CD')
    ok = fail = 0
    for pr in project_rows:
        r = pr['idx']
        m = matched.get(r)
        if m:
            addr = ws.cell(row=r, column=col_addr, value=str(m['page']['id']))
            addr.number_format = '@'          # 텍스트 — 엑셀이 .0/지수로 바꾸지 않게
            ws.cell(row=r, column=col_result, value='성공')
            ws.cell(row=r, column=col_title, value=m['page']['title'])
            ws.cell(row=r, column=col_score, value=round(m['title_sim'], 2))
            ok += 1
        else:
            cell = ws.cell(row=r, column=col_result, value='실패')   # 컨플 주소는 기존 값 그대로
            cell.fill = fail_fill
            ws.cell(row=r, column=col_title, value=None)
            ws.cell(row=r, column=col_score, value=None)
            fail += 1
    for col, width in ((col_result, 10), (col_title, 40), (col_score, 8)):
        ws.column_dimensions[get_column_letter(col)].width = width

    # 미매칭 페이지 시트(과제 단계 후보 중 어느 행에도 쓰이지 않은 것)
    used_ids = {m['page']['id'] for m in matched.values()}
    leftovers = [p for p in pages if p['id'] not in used_ids]
    name = '미매칭 페이지'
    if name in wb.sheetnames:
        del wb[name]
    ws2 = wb.create_sheet(name)
    ws2.append(['제목', '페이지 ID', '부서(상위 페이지 제목)', '가장 비슷한 과제명', '유사도'])
    for c in ws2[1]:
        c.font = bold
    for p in leftovers:
        best, best_sim = '', 0.0
        for pr in project_rows:
            s = similarity(pr['project'], p['title'])
            if s > best_sim:
                best, best_sim = pr['project'], s
        row = [p['title'], str(p['id']), p.get('parent_title', ''), best, round(best_sim, 2) if best else None]
        ws2.append(row)
        ws2.cell(row=ws2.max_row, column=2).number_format = '@'
    for i, w in enumerate((40, 14, 30, 40, 8), 1):
        ws2.column_dimensions[get_column_letter(i)].width = w
    return {'success': ok, 'fail': fail, 'unmatched_pages': len(leftovers), 'candidates': len(pages),
            'result_column': get_column_letter(col_result)}


def build_updated_workbook(rows: list[dict]) -> tuple[bytes, dict, str]:
    """(xlsx 바이트, 요약, 출처 설명). 필요한 입력이 없으면 ValueError."""
    if not rows:
        raise ValueError('추출 결과가 없습니다 — 먼저 하위 페이지를 추출하세요.')
    wb, source = _load_workbook_from_source()
    summary = apply_to_workbook(wb, rows)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), summary, source
