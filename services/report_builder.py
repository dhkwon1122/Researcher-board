"""
업무 시나리오 > 조서 작성(2026-10-15, 사용자 요청) — 엑셀 열을 사용자가 직접 구성하는 다운로드.

  - 기본 항목(FIELDS): 지금 연구원 명단 엑셀에서 "성명(성별/나이)"처럼 묶여 있던 열을 쪼갠 단위 항목.
    값 계산은 services/researcher_profile_export의 행 컨텍스트·열 함수를 그대로 재사용한다(같은 데이터·같은 표기).
  - 양식(template): 열 목록 [{header, fields:[항목키...], join}] + max_items(여러 건 항목은 최근 N건만).
    join(묶는 방식): 정해진 형식 중 선택(사용자 확정 — 자유 서식 문자열은 쓰지 않음).
  - 저장: data/processed/report_templates.json — 개인 양식(owner=사용자ID) + 공용 양식(shared, 관리자만 저장·삭제).
    "기본 양식"(코드 내장)은 현재 연구원 명단 엑셀과 같은 구성.
  - 권한: 평가 항목은 view_evaluation, 핵심이력은 view_incentive 권한자에게만 목록에 보이고, 엑셀을 만들 때도 다시 확인한다.
기존 연구원 명단 엑셀 다운로드(pages/researcher_list.py)는 그대로 둔다(사용자 확정).
"""
from __future__ import annotations

import io
import json
import os
import threading
import uuid
from datetime import date, datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from services import data_store
from services import researcher_profile_export as rpe
from services.excel_text import excel_safe

TEMPLATES_PATH = os.path.join(data_store.DATA_DIR, 'report_templates.json')
_LOCK = threading.Lock()
DEFAULT_TEMPLATE_ID = 'builtin-default'


# ─── 기본 항목 ───────────────────────────────────────────────────────────────

def _r(rows) -> dict:
    return rows.get('researcher') or {}


def _name(_rid, rows):
    return rpe._or_dash(_r(rows).get('name'))


def _gender(_rid, rows):
    return rpe._s(_r(rows).get('gender')) or '-'


def _age(_rid, rows):
    by = rpe._birth_year_int(_r(rows).get('birth_year'))
    return f'{datetime.now().year - by}세' if by is not None else '-'


def _birth_year(_rid, rows):
    by = rpe._birth_year_int(_r(rows).get('birth_year'))
    return str(by) if by is not None else '-'


def _dept(_rid, rows):
    org = rpe._s(_r(rows).get('org_code'))
    dep_map, _pjt = rows.get('dep_pjt_maps') or ({}, {})
    return rpe._or_dash(dep_map.get(org) or _r(rows).get('department'))


def _project(_rid, rows):
    org = rpe._s(_r(rows).get('org_code'))
    _dep, pjt_map = rows.get('dep_pjt_maps') or ({}, {})
    return rpe._or_dash(pjt_map.get(org) or org)


def _hire_date(_rid, rows):
    hd = rpe._s(_r(rows).get('hire_date'))
    return hd[:10].replace('-', '') if hd else '-'


def _tenure(_rid, rows):
    hd = rpe._s(_r(rows).get('hire_date'))
    try:
        return f"{round((date.today() - date.fromisoformat(hd[:10])).days / 365, 1)}년"
    except ValueError:
        return '-'


def _cl(_rid, rows):
    return rpe._or_dash(_r(rows).get('position'))


def _cl_years(_rid, rows):
    y = rpe.position_years(_r(rows).get('promotion_date'))
    return f'{y}년차' if y is not None else '-'


def _highest(key: str):
    def _fn(_rid, rows):
        row = rpe.highest_degree_row(rows.get('education') or [])
        if not row:
            return '-'
        if key == 'major' and rpe._s(row.get('degree')) == '고교':
            return '-'
        return rpe._or_dash(row.get(key))
    return _fn


# (키, 표시명, 그룹, 함수, 필요 권한 또는 None, 여러 건 항목 여부)
def _field_defs() -> list:
    defs = [
        ('researcher_id', '사번', '인적', rpe._col_id, None, False),
        ('knox_id', 'Knox ID', '인적', rpe._col_knox, None, False),
        ('name', '성명', '인적', _name, None, False),
        ('gender', '성별', '인적', _gender, None, False),
        ('age', '나이', '인적', _age, None, False),
        ('birth_year', '출생연도', '인적', _birth_year, None, False),
        ('org_code', '조직코드', '소속', rpe._col_org_code, None, False),
        ('dept', '부서(1단계)', '소속', _dept, None, False),
        ('project', '과제(3단계)', '소속', _project, None, False),
        ('position_title', '직책', '소속', rpe._col_position_title, None, False),
        ('job_function', '직무', '소속', rpe._col_job_function, None, False),
        ('employment_status', '재직상태', '소속', rpe._col_employment_status, None, False),
        ('hire_date', '입사일', '경력', _hire_date, None, False),
        ('tenure', '근속년수', '경력', _tenure, None, False),
        ('cl', 'CL', '경력', _cl, None, False),
        ('cl_years', '년차', '경력', _cl_years, None, False),
        ('cl_and_years', 'CL/년차', '경력', rpe._col_position_year, None, False),
        ('degree', '최종학력', '학력', _highest('degree'), None, False),
        ('school', '최종학교', '학력', _highest('school'), None, False),
        ('major', '최종전공', '학력', _highest('major'), None, False),
        ('education_all', '학력(전체)', '학력', rpe._col_education, None, True),
        ('tasks', '과제수행이력', '실적·이력', rpe._col_tasks, None, True),
        ('publications', '논문 실적', '실적·이력', rpe._col_publications, None, True),
        ('patents', '특허 실적', '실적·이력', rpe._col_patents, None, True),
        ('language', '어학', '실적·이력', rpe._col_language, None, True),
        ('tech_ownership', '보유기술', '실적·이력', rpe._col_tech_ownership, None, True),
        ('job_profile', '직무이력', '실적·이력', rpe._col_job_profile, None, True),
        ('work_experience', '근무 경력', '실적·이력', rpe._col_work_experience, None, True),
        ('nurturing', '양성이력', '실적·이력', rpe._col_nurturing, None, True),
        ('incentive', '핵심이력', '실적·이력', rpe._col_incentive, 'view_incentive', True),
    ]
    defs += [(f'expertise_{i}', label.replace('보유전문성(', '보유전문성 - ').rstrip(')'), '보유 전문성(AI)', fn, None, True)
             for i, (label, fn) in enumerate(rpe._EXPERTISE_COLUMNS)]
    defs.append(('eval_summary', rpe._EVAL_HEADER.replace('\n', ' '), '평가', rpe._col_evaluation, 'view_evaluation', False))
    defs += [(f'eval_salary_{i}', label, '평가', fn, 'view_evaluation', False)
             for i, (label, fn) in enumerate(rpe._EVAL_SALARY_COLUMNS)]
    defs += [(f'eval_half_{i}', label, '평가', fn, 'view_evaluation', False)
             for i, (label, fn) in enumerate(rpe._EVAL_HALF_COLUMNS)]
    return defs


FIELDS = {k: {'label': lbl, 'group': grp, 'fn': fn, 'perm': perm, 'multi': multi}
          for k, lbl, grp, fn, perm, multi in _field_defs()}
GROUP_ORDER = ['인적', '소속', '경력', '학력', '실적·이력', '보유 전문성(AI)', '평가']


def field_options(permissions: dict) -> list:
    """드롭다운 옵션(권한 없는 항목 제외) — 라벨에 그룹을 붙여 찾기 쉽게."""
    out = []
    for g in GROUP_ORDER:
        for k, f in FIELDS.items():
            if f['group'] == g and (not f['perm'] or permissions.get(f['perm'])):
                out.append({'label': f"[{g}] {f['label']}", 'value': k})
    return out


# ─── 묶는 방식 ───────────────────────────────────────────────────────────────

JOINS = {
    'paren': '괄호: A (B/C)',
    'paren_newline': '줄바꿈 괄호: A↵(B/C)',
    'slash': '슬래시: A/B/C',
    'comma': '쉼표: A, B, C',
    'space': '공백: A B C',
    'newline': '줄바꿈: A↵B↵C',
}


def join_values(values: list, how: str) -> str:
    vals = [str(v) for v in values]
    if vals and all(v.strip() in ('', '-') for v in vals):   # 전부 비어 있으면 "-" 하나로
        return '-'
    if len(vals) <= 1:
        return vals[0] if vals else '-'
    if how == 'paren':
        return f"{vals[0]} ({'/'.join(vals[1:])})"
    if how == 'paren_newline':
        return f"{vals[0]}\n({'/'.join(vals[1:])})"
    return {'slash': '/', 'comma': ', ', 'space': ' ', 'newline': '\n'}.get(how, '/').join(vals)


def _limit(value: str, max_items: int | None) -> str:
    if not max_items or not value or value == '-':
        return value
    lines = value.split('\n')
    if len(lines) <= max_items:
        return value
    return '\n'.join(lines[:max_items] + [f'외 {len(lines) - max_items}건'])


# ─── 양식 저장소 ─────────────────────────────────────────────────────────────

def default_template() -> dict:
    """현재 연구원 명단 엑셀(기본 열)과 같은 구성."""
    cols = [
        {'header': '조직코드', 'fields': ['org_code'], 'join': 'slash'},
        {'header': '사번', 'fields': ['researcher_id'], 'join': 'slash'},
        {'header': 'Knox ID', 'fields': ['knox_id'], 'join': 'slash'},
        {'header': '성명\n(성별/나이)', 'fields': ['name', 'gender', 'age'], 'join': 'paren_newline'},
        {'header': '부서\n(과제)', 'fields': ['dept', 'project'], 'join': 'paren_newline'},
        {'header': '입사일', 'fields': ['hire_date', 'tenure'], 'join': 'paren_newline'},
        {'header': '학력', 'fields': ['education_all'], 'join': 'slash'},
        {'header': 'CL/년차', 'fields': ['cl_and_years'], 'join': 'slash'},
        {'header': '직책', 'fields': ['position_title'], 'join': 'slash'},
    ]
    return {'id': DEFAULT_TEMPLATE_ID, 'name': '기본 양식(연구원 명단과 동일)', 'owner': '', 'shared': True,
            'builtin': True, 'columns': cols, 'max_items': None}


def _read_all() -> list:
    try:
        with open(TEMPLATES_PATH, encoding='utf-8') as f:
            data = json.load(f)
        return data if isinstance(data, list) else []
    except (OSError, ValueError):
        return []


def _write_all(items: list) -> None:
    os.makedirs(os.path.dirname(TEMPLATES_PATH), exist_ok=True)
    tmp = TEMPLATES_PATH + '.tmp'
    with open(tmp, 'w', encoding='utf-8') as f:
        json.dump(items, f, ensure_ascii=False, indent=2)
    os.replace(tmp, TEMPLATES_PATH)


def list_templates(user_id: str) -> list:
    """기본 양식 + 공용 양식 + 내 양식."""
    items = [t for t in _read_all() if t.get('shared') or t.get('owner') == user_id]
    items.sort(key=lambda t: (not t.get('shared'), t.get('name', '')))
    return [default_template()] + items


def get_template(template_id: str, user_id: str) -> dict | None:
    return next((t for t in list_templates(user_id) if t['id'] == template_id), None)


def clean_columns(columns: list) -> list:
    out = []
    for c in columns or []:
        fields = [f for f in (c.get('fields') or []) if f in FIELDS]
        if not fields:
            continue
        header = str(c.get('header') or '').strip() or join_values([FIELDS[f]['label'] for f in fields], c.get('join'))
        out.append({'header': header, 'fields': fields, 'join': c.get('join') if c.get('join') in JOINS else 'slash'})
    return out


def save_template(user_id: str, name: str, columns: list, max_items, shared: bool = False,
                  is_admin: bool = False, template_id: str | None = None) -> dict:
    """새로 저장하거나(template_id 없음) 덮어쓴다. 공용은 관리자만. 반환: 저장된 양식."""
    name = (name or '').strip()
    if not name:
        raise ValueError('양식 이름을 입력하세요.')
    columns = clean_columns(columns)
    if not columns:
        raise ValueError('항목이 들어 있는 열이 하나 이상 있어야 합니다.')
    if shared and not is_admin:
        raise PermissionError('공용 양식은 관리자만 저장할 수 있습니다.')
    with _LOCK:
        items = _read_all()
        existing = next((t for t in items if t['id'] == template_id), None) if template_id else None
        if existing is not None:
            if existing.get('shared') and not is_admin:
                raise PermissionError('공용 양식은 관리자만 수정할 수 있습니다.')
            if not existing.get('shared') and existing.get('owner') != user_id:
                raise PermissionError('다른 사람의 양식은 수정할 수 없습니다.')
            items.remove(existing)
        tpl = {'id': existing['id'] if existing else uuid.uuid4().hex[:12], 'name': name, 'owner': user_id,
               'shared': bool(shared), 'columns': columns,
               'max_items': int(max_items) if str(max_items or '').strip().isdigit() and int(max_items) > 0 else None,
               'updated_at': datetime.now().strftime('%Y-%m-%d %H:%M')}
        items.append(tpl)
        _write_all(items)
    return tpl


def delete_template(template_id: str, user_id: str, is_admin: bool = False) -> None:
    with _LOCK:
        items = _read_all()
        t = next((x for x in items if x['id'] == template_id), None)
        if t is None:
            raise ValueError('삭제할 양식을 찾을 수 없습니다(기본 양식은 삭제할 수 없음).')
        if (t.get('shared') and not is_admin) or (not t.get('shared') and t.get('owner') != user_id):
            raise PermissionError('삭제 권한이 없습니다.')
        items.remove(t)
        _write_all(items)


# ─── 대상자 ──────────────────────────────────────────────────────────────────

def resolve_targets(platforms=None, projects=None, researcher_ids=None) -> list:
    """현재 재직자 중 선택한 플랫폼(1단계)·과제(3단계) 소속 + 개별 지정 인원(합집합)."""
    from services import similarity_map as sm
    res = data_store.filter_current(data_store.read_processed('researchers'))
    out = set(str(r).zfill(8) for r in researcher_ids or [])
    platforms, projects = set(platforms or []), set(projects or [])
    if (platforms or projects) and not res.empty and 'org_code' in res.columns:
        cache: dict = {}
        for rid, org in zip(res['researcher_id'], res['org_code'].astype(str).str.strip()):
            if org not in cache:
                cache[org] = sm.org_code_level_names(org)
            l1, _l2, l3 = cache[org]
            if (l1 and l1 in platforms) or (l3 and l3 in projects):
                out.add(rid)
    return sorted(out)


# ─── 값 계산 / 엑셀 ──────────────────────────────────────────────────────────

def _allowed(field: str, permissions: dict) -> bool:
    f = FIELDS.get(field)
    return bool(f) and (not f['perm'] or permissions.get(f['perm']))


def build_table(researcher_ids: list, columns: list, max_items=None) -> tuple[list, list, list]:
    """(헤더 목록, 행 목록, 실제 열 정의). 권한 없는 항목은 빼고(열의 항목이 전부 빠지면 열도 생략) 계산."""
    ids, ctx_for, permissions = rpe.prepare_rows(list(researcher_ids))
    cols = []
    for c in clean_columns(columns):
        fields = [f for f in c['fields'] if _allowed(f, permissions)]
        if fields:
            cols.append({**c, 'fields': fields})
    rows = []
    for rid in ids:
        ctx = ctx_for(rid)
        rows.append([join_values([_limit(str(FIELDS[f]['fn'](rid, ctx)), max_items if FIELDS[f]['multi'] else None)
                                  for f in c['fields']], c['join']) for c in cols])
    return [c['header'] for c in cols], rows, cols


def build_workbook(researcher_ids: list, columns: list, max_items=None, sheet_title: str = '조서') -> bytes:
    headers, rows, cols = build_table(researcher_ids, columns, max_items)
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31] or '조서'
    head_font = Font(name=rpe._FONT_NAME, size=rpe._FONT_SIZE, bold=True)
    body_font = Font(name=rpe._FONT_NAME, size=rpe._FONT_SIZE)
    center = Alignment(wrap_text=True, vertical='center', horizontal='center')
    left = Alignment(wrap_text=True, vertical='center', horizontal='left')
    for j, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=j, value=h)
        cell.font, cell.border, cell.alignment = head_font, rpe._BORDER, center
    multi_col = [any(FIELDS[f]['multi'] for f in c['fields']) for c in cols]
    for i, row in enumerate(rows, start=2):
        for j, v in enumerate(row, start=1):
            cell = ws.cell(row=i, column=j, value=excel_safe(v))
            cell.font, cell.border = body_font, rpe._BORDER
            cell.alignment = left if multi_col[j - 1] else center
    for j, h in enumerate(headers, start=1):
        longest = max([len(line) for v in [h] + [r[j - 1] for r in rows[:200]] for line in str(v).split('\n')] or [8])
        ws.column_dimensions[get_column_letter(j)].width = max(8, min(60, longest * 1.6 + 2))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def default_filename(name: str = '') -> str:
    stem = ''.join(ch for ch in (name or '조서') if ch not in '\\/:*?"<>|').strip() or '조서'
    return f"{stem}_{datetime.now().strftime('%Y%m%d%H%M')}.xlsx"
